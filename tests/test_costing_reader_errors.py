"""Đọc file Costing người dùng sửa: mọi lỗi phải chỉ đúng ô trong Excel.

CLAUDE.md: workbook có đúng hai sheet `Hướng dẫn` và `Costing`; lỗi file là lỗi
người dùng nên phải có mã ổn định và chỉ rõ chỗ cần sửa, không được thoát ra
thành lỗi hệ thống.
"""

from __future__ import annotations

import pytest
from openpyxl import Workbook, load_workbook

from tests.test_costing_workbook import sample_document
from wfx_panel.workbooks import costing as costing_workbook
from wfx_panel.workbooks.costing.reader import _worksheet_rows

GUIDE = costing_workbook.GUIDE_SHEET
FORM = costing_workbook.FORM_SHEET


@pytest.fixture
def workbook_path(tmp_path):
    target = tmp_path / "SWN0000001-Costing.xlsx"
    costing_workbook.write_costing_file(sample_document(), target)
    return target


def _edit(path):
    workbook = load_workbook(path)
    return workbook, workbook[FORM], workbook[GUIDE]


def _headers(form):
    return {
        str(cell.value): cell.column
        for cell in form[1]
        if cell.value is not None
    }


def _error(path):
    with pytest.raises(costing_workbook.CostingWorkbookError) as raised:
        costing_workbook.read_costing_xlsx(path)
    return raised.value


# --- sheet và cột --------------------------------------------------------


def test_an_extra_sheet_in_the_workbook_is_refused(workbook_path):
    workbook, _form, _guide = _edit(workbook_path)
    workbook.create_sheet("Ghi chú")
    workbook.save(workbook_path)

    assert _error(workbook_path).code == "COSTING_FORMAT_UNSUPPORTED"


def test_a_form_sheet_missing_a_standard_column_is_refused(workbook_path):
    workbook, form, _guide = _edit(workbook_path)
    form.cell(1, 1, "Khac")
    workbook.save(workbook_path)

    error = _error(workbook_path)

    assert error.code == "COSTING_FORMAT_UNSUPPORTED"
    assert any("thiếu cột" in detail for detail in error.details)


def test_a_workbook_without_a_style_code_is_refused(workbook_path):
    workbook, _form, guide = _edit(workbook_path)
    guide["B3"] = ""
    workbook.save(workbook_path)

    error = _error(workbook_path)

    assert error.code == "COSTING_VALIDATION_FAILED"
    assert "Style Code" in error.details[0]


def test_a_file_that_is_not_a_workbook_at_all_is_refused(tmp_path):
    broken = tmp_path / "hong.xlsx"
    broken.write_text("khong phai xlsx", encoding="utf-8")

    assert _error(broken).code == "COSTING_VALIDATION_FAILED"


def test_a_zip_openpyxl_cannot_read_is_reported_as_an_unreadable_workbook(
    tmp_path,
):
    import zipfile

    target = tmp_path / "gia.xlsx"
    with zipfile.ZipFile(target, "w") as archive:
        archive.writestr("readme.txt", b"khong phai workbook")

    error = _error(target)

    assert error.code == "COSTING_VALIDATION_FAILED"
    assert "Không đọc được workbook" in error.message


# --- từng dòng trong sheet Costing ---------------------------------------


def _first_item_row(form, headers):
    for row in range(2, form.max_row + 1):
        if str(form.cell(row, headers["__Item Key"]).value or "").strip():
            return row
    raise AssertionError("Không tìm thấy dòng Article trong form")


def test_an_action_the_format_does_not_define_is_reported_with_its_cell(
    workbook_path,
):
    workbook, form, _guide = _edit(workbook_path)
    headers = _headers(form)
    row = _first_item_row(form, headers)
    form.cell(row, headers["Action"], "REMOVE")
    workbook.save(workbook_path)

    error = _error(workbook_path)

    assert error.code == "COSTING_VALIDATION_FAILED"
    assert any(f"{FORM}!B{row}" in detail for detail in error.details)


def test_an_item_type_the_format_does_not_define_is_reported_with_its_cell(
    workbook_path,
):
    workbook, form, _guide = _edit(workbook_path)
    headers = _headers(form)
    row = _first_item_row(form, headers)
    form.cell(row, headers["__Item Type"], "material")
    workbook.save(workbook_path)

    error = _error(workbook_path)

    assert any("Item Type" in detail for detail in error.details)


def test_two_rows_sharing_an_item_key_are_reported_with_the_first_row(
    workbook_path,
):
    workbook, form, _guide = _edit(workbook_path)
    headers = _headers(form)
    row = _first_item_row(form, headers)
    target_row = form.max_row + 1
    for column in headers.values():
        form.cell(target_row, column, form.cell(row, column).value)
    workbook.save(workbook_path)

    error = _error(workbook_path)

    assert any(
        "Item Key trùng dòng" in detail and str(row) in detail
        for detail in error.details
    )


# --- đọc sheet thô --------------------------------------------------------


def test_a_sheet_with_no_header_row_at_all_yields_no_rows():
    workbook = Workbook()
    sheet = workbook.active

    assert list(_worksheet_rows(sheet)) == []


def test_rows_that_are_completely_empty_are_skipped():
    workbook = Workbook()
    sheet = workbook.active
    sheet.append(["A", "B"])
    sheet.append([None, None])
    sheet.append(["x", "y"])

    rows = list(_worksheet_rows(sheet))

    assert [row["A"] for row in rows] == ["x"]


def test_columns_the_reader_is_told_to_ignore_are_left_out():
    workbook = Workbook()
    sheet = workbook.active
    sheet.append(["A", "B"])
    sheet.append(["x", "y"])

    rows = list(_worksheet_rows(sheet, ignored_columns={"B"}))

    assert rows[0] == {"A": "x", "__Excel Row": 2}
