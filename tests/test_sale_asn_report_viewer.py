"""Report Viewer của Sale ASN Documents: chờ, tải, nhận diện và dọn popup.

`tests/test_sale_asn_documents.py` đã phủ đường đi thành công. File này phủ các
nhánh mà một phiên WFX thật vẫn gặp: report cũ còn trên màn hình, SSRS chưa gán
`ExportUrlBase`, popup bị đóng giữa chừng và file tải về không phải Excel.
"""

from __future__ import annotations

from io import BytesIO

import pytest
from openpyxl import Workbook

import wfx_panel.automation.sale_asn_documents.report as report
from tests.fakes.mini_dom import Element, MiniFrame
from tests.fakes.module_reflection import patch_automation
from tests.fakes.wfx_dom import install_fake_clock
from wfx_panel.automation._common import PlaywrightError, PlaywrightTimeoutError
from wfx_panel.automation.sale_asn_documents.constants import (
    BUYER_INVOICE_SELECTOR,
    PACKING_LIST_SELECTOR,
)


class ReportFrame(MiniFrame):
    """MiniFrame cộng `goto` và các lỗi điều hướng mà Playwright có thể ném."""

    def __init__(self, *args, goto_error=None, **kwargs) -> None:
        super().__init__(*args, **kwargs)
        self.goto_error = goto_error
        self.goto_calls: list[str] = []

    def goto(self, url, **_kwargs):
        self.goto_calls.append(url)
        if self.goto_error is not None:
            raise self.goto_error


class BrokenFrame:
    """Frame đang bị detach: mọi truy vấn selector đều ném PlaywrightError."""

    url = "https://wfx.test/detached"

    def locator(self, _selector):
        raise PlaywrightError("frame đã detach")

    def evaluate(self, _script, _arg=None):
        raise PlaywrightError("frame đã detach")


class FakePage:
    def __init__(self, *frames, url="https://wfx.test/sale-asn", clock=None):
        self.frames = list(frames)
        self.url = url
        self.clock = clock
        self.closed = False
        self.close_calls = 0
        self.close_error: BaseException | None = None
        self.go_back_calls = 0
        self.go_back_error: BaseException | None = None
        for frame in self.frames:
            if isinstance(frame, MiniFrame):
                frame.page = self

    def wait_for_timeout(self, milliseconds):
        if self.clock is not None:
            self.clock.advance(float(milliseconds) / 1_000.0)

    def is_closed(self):
        return self.closed

    def close(self, run_before_unload=True):
        self.close_calls += 1
        if self.close_error is not None:
            raise self.close_error
        self.closed = True

    def go_back(self, **_kwargs):
        self.go_back_calls += 1
        if self.go_back_error is not None:
            raise self.go_back_error


class FakeContext:
    def __init__(self, *pages):
        self.pages = list(pages)


def _export_node(*, enabled: bool = True, visible: bool = True) -> Element:
    return Element(
        "a",
        attrs={"title": "Export drop down menu"},
        enabled=enabled,
        visible=visible,
    )


def _report_frame(
    *,
    clock=None,
    enabled: bool = True,
    visible: bool = True,
    loading: bool = False,
    export_url: object = "/Reserved.ReportViewerWebControl.axd?Format=",
    url: str = "https://wfx.test/report/Viewer.aspx",
) -> ReportFrame:
    children = [_export_node(enabled=enabled, visible=visible)]
    if loading:
        children.append(
            Element(
                "div",
                id="ctl00_AsyncWait",
                attrs={"style": "display: block"},
            )
        )
    scripts = {
        "ExportUrlBase": (
            export_url if callable(export_url) else (lambda _arg: export_url)
        )
    }
    return ReportFrame(
        Element("body", children=children),
        url=url,
        clock=clock,
        scripts=scripts,
    )


def _documents_frame_node(clock=None) -> ReportFrame:
    return ReportFrame(
        Element(
            "body",
            children=[
                Element("a", id=PACKING_LIST_SELECTOR.lstrip("#")),
                Element("a", id=BUYER_INVOICE_SELECTOR.lstrip("#")),
            ],
        ),
        url="https://wfx.test/SaleASNDocuments.aspx",
        clock=clock,
    )


# --- tìm frame ----------------------------------------------------------


def test_finding_a_frame_skips_detached_frames_and_prefers_the_newest_page(
    monkeypatch,
):
    clock = install_fake_clock(monkeypatch, report)
    stale = _documents_frame_node(clock)
    fresh = _documents_frame_node(clock)
    old_page = FakePage(BrokenFrame(), stale, clock=clock)
    new_page = FakePage(fresh, clock=clock)
    context = FakeContext(old_page, new_page)

    page, frame = report._find_frame_with(
        context,
        (PACKING_LIST_SELECTOR, BUYER_INVOICE_SELECTOR),
        timeout_s=5,
        visible=True,
    )

    assert page is new_page
    assert frame is fresh


def test_finding_a_frame_names_every_selector_it_waited_for(monkeypatch):
    clock = install_fake_clock(monkeypatch, report)
    context = FakeContext(FakePage(BrokenFrame(), clock=clock))

    with pytest.raises(PlaywrightTimeoutError) as error:
        report._find_frame_with(
            context,
            (PACKING_LIST_SELECTOR, BUYER_INVOICE_SELECTOR),
            timeout_s=1,
        )

    assert PACKING_LIST_SELECTOR in str(error.value)
    assert BUYER_INVOICE_SELECTOR in str(error.value)


def test_a_frame_that_only_has_one_of_the_two_document_links_is_not_enough(
    monkeypatch,
):
    clock = install_fake_clock(monkeypatch, report)
    half = ReportFrame(
        Element(
            "body",
            children=[Element("a", id=PACKING_LIST_SELECTOR.lstrip("#"))],
        ),
        clock=clock,
    )
    context = FakeContext(FakePage(half, clock=clock))

    with pytest.raises(PlaywrightTimeoutError):
        report._documents_frame(context, timeout_s=1)


def test_documents_frame_ignores_links_that_are_present_but_hidden(monkeypatch):
    clock = install_fake_clock(monkeypatch, report)
    hidden = ReportFrame(
        Element(
            "body",
            children=[
                Element(
                    "a", id=PACKING_LIST_SELECTOR.lstrip("#"), visible=False
                ),
                Element("a", id=BUYER_INVOICE_SELECTOR.lstrip("#")),
            ],
        ),
        clock=clock,
    )
    context = FakeContext(FakePage(hidden, clock=clock))

    with pytest.raises(PlaywrightTimeoutError):
        report._documents_frame(context, timeout_s=1)


def test_documents_frame_returns_the_page_that_owns_the_links(monkeypatch):
    clock = install_fake_clock(monkeypatch, report)
    docs = _documents_frame_node(clock)
    page = FakePage(docs, clock=clock)

    found_page, found_frame = report._documents_frame(
        FakeContext(page), timeout_s=5
    )

    assert (found_page, found_frame) == (page, docs)


# --- đánh dấu report cũ -------------------------------------------------


def test_marking_report_frames_only_touches_frames_that_can_export(monkeypatch):
    install_fake_clock(monkeypatch, report)
    exporting = _report_frame()
    plain = _documents_frame_node()
    context = FakeContext(FakePage(exporting, plain, BrokenFrame()))

    snapshots = report._mark_report_frames(context)

    assert [frame for frame, _marker in snapshots] == [exporting]
    assert snapshots[0][1].startswith("asn-report-")
    assert exporting.markers["__wfxAsnReportMarker"] == snapshots[0][1]


def test_a_frame_missing_from_the_snapshot_counts_as_a_brand_new_report():
    assert report._report_frame_is_new(_report_frame(), []) is True


def test_a_frame_keeping_its_marker_is_the_same_report_as_before(monkeypatch):
    install_fake_clock(monkeypatch, report)
    frame = _report_frame()
    snapshots = report._mark_report_frames(FakeContext(FakePage(frame)))

    assert report._report_frame_is_new(frame, snapshots) is False


def test_a_reloaded_frame_loses_its_marker_and_counts_as_a_new_report(
    monkeypatch,
):
    install_fake_clock(monkeypatch, report)
    frame = _report_frame()
    snapshots = report._mark_report_frames(FakeContext(FakePage(frame)))
    frame.markers.clear()

    assert report._report_frame_is_new(frame, snapshots) is True


def test_a_frame_that_detaches_between_marking_and_reading_counts_as_new():
    broken = BrokenFrame()

    assert report._report_frame_is_new(broken, [(broken, "asn-report-1")]) is True


def test_export_url_is_empty_when_the_viewer_frame_is_already_gone():
    def explode(_arg):
        raise PlaywrightError("Execution context was destroyed")

    frame = _report_frame(export_url=explode)

    assert report._report_export_url(frame) == ""


def test_export_url_is_empty_while_ssrs_has_not_assigned_it_yet():
    assert report._report_export_url(_report_frame(export_url=None)) == ""


# --- chờ Report Viewer --------------------------------------------------


def test_waiting_for_the_report_ignores_the_report_left_over_from_last_run(
    monkeypatch,
):
    clock = install_fake_clock(monkeypatch, report)
    stale = _report_frame(clock=clock)
    page = FakePage(stale, clock=clock)
    context = FakeContext(page)
    snapshots = report._mark_report_frames(context)

    with pytest.raises(PlaywrightTimeoutError, match="chưa load xong"):
        report._wait_report_ready(context, snapshots, timeout_s=3)


def test_waiting_for_the_report_returns_the_frame_only_after_it_stays_ready(
    monkeypatch,
):
    clock = install_fake_clock(monkeypatch, report)
    fresh = _report_frame(clock=clock)
    page = FakePage(fresh, clock=clock)
    context = FakeContext(page)
    started = clock.monotonic()

    found_page, found_frame = report._wait_report_ready(
        context, [], timeout_s=30
    )

    assert (found_page, found_frame) == (page, fresh)
    # Toolbar sáng ngay từ lượt quét đầu, nhưng hàm vẫn phải giữ đủ 0,8 giây
    # ổn định trước khi cho phép click Export.
    assert clock.monotonic() - started >= 0.8


def test_a_toolbar_that_lights_up_before_ssrs_binds_the_export_url_is_not_ready(
    monkeypatch,
):
    clock = install_fake_clock(monkeypatch, report)
    urls = ["", "", "/Reserved.ReportViewerWebControl.axd?Format="]

    def export_url(_arg):
        return urls.pop(0) if len(urls) > 1 else urls[0]

    fresh = _report_frame(clock=clock, export_url=export_url)
    page = FakePage(fresh, clock=clock)

    _found_page, found_frame = report._wait_report_ready(
        FakeContext(page), [], timeout_s=30
    )

    assert found_frame is fresh
    assert urls == ["/Reserved.ReportViewerWebControl.axd?Format="]


def test_a_report_still_showing_its_async_spinner_is_not_ready(monkeypatch):
    clock = install_fake_clock(monkeypatch, report)
    loading = _report_frame(clock=clock, loading=True)

    with pytest.raises(PlaywrightTimeoutError, match="chưa load xong"):
        report._wait_report_ready(
            FakeContext(FakePage(loading, clock=clock)), [], timeout_s=3
        )


def test_a_disabled_export_button_never_settles_into_ready(monkeypatch):
    clock = install_fake_clock(monkeypatch, report)
    disabled = _report_frame(clock=clock, enabled=False)

    with pytest.raises(PlaywrightTimeoutError, match="chưa load xong"):
        report._wait_report_ready(
            FakeContext(FakePage(disabled, clock=clock)), [], timeout_s=3
        )


def test_waiting_for_a_report_that_never_opens_reports_the_viewer_not_the_frame(
    monkeypatch,
):
    clock = install_fake_clock(monkeypatch, report)
    context = FakeContext(FakePage(BrokenFrame(), clock=clock))

    with pytest.raises(PlaywrightTimeoutError, match="chưa load xong"):
        report._wait_report_ready(context, [], timeout_s=2)


# --- tải file -----------------------------------------------------------


class FetchFrame:
    """Frame chỉ mô phỏng vòng fetch trong Report Viewer."""

    url = "https://wfx.test/report/Viewer.aspx"

    def __init__(self, states, *, clock, chunks=b"", cleanup_error=None):
        self.states = list(states)
        self.clock = clock
        self.chunks = chunks
        self.cleanup_error = cleanup_error
        self.starts: list[str] = []
        self.cleanups = 0

    def wait_for_timeout(self, milliseconds):
        self.clock.advance(float(milliseconds) / 1_000.0)

    def evaluate(self, script, argument=None):
        if script == report._REPORT_FETCH_START_JS:
            self.starts.append(argument)
            return True
        if script == report._REPORT_FETCH_STATE_JS:
            if len(self.states) > 1:
                return self.states.pop(0)
            return self.states[0]
        if script == report._REPORT_FETCH_CHUNK_JS:
            import base64

            window = self.chunks[
                argument["offset"] : argument["offset"] + argument["size"]
            ]
            return base64.b64encode(window).decode("ascii")
        if script == report._REPORT_FETCH_CLEANUP_JS:
            self.cleanups += 1
            if self.cleanup_error is not None:
                raise self.cleanup_error
            return None
        raise AssertionError(f"script lạ: {script[:60]}")


def _workbook_bytes(title: str) -> bytes:
    workbook = Workbook()
    workbook.active["A1"] = title
    buffer = BytesIO()
    workbook.save(buffer)
    workbook.close()
    return buffer.getvalue()


def test_download_refuses_to_start_before_wfx_created_the_export_link(
    tmp_path, monkeypatch
):
    clock = install_fake_clock(monkeypatch, report)
    patch_automation(monkeypatch, report, "_report_export_url", lambda _f: "")
    frame = FetchFrame([{"done": True}], clock=clock)

    with pytest.raises(PlaywrightTimeoutError, match="chưa tạo link download"):
        report._download_report_excel(
            object(),
            frame,
            tmp_path / "packing.xlsx",
            "Packing List",
            lambda _message: None,
        )

    assert frame.starts == []


def test_a_slow_report_keeps_telling_the_user_wfx_is_still_generating_it(
    tmp_path, monkeypatch
):
    clock = install_fake_clock(monkeypatch, report)
    patch_automation(
        monkeypatch, report, "_report_export_url", lambda _f: "/export?Format="
    )
    payload = _workbook_bytes("PACKING LIST")
    pending = {"done": False}
    finished = {
        "done": True,
        "ok": True,
        "status": 200,
        "size": len(payload),
        "prefix": "PK",
    }
    frame = FetchFrame(
        [pending] * 400 + [finished], clock=clock, chunks=payload
    )
    logs: list[str] = []
    target = tmp_path / "packing.xlsx"

    report._download_report_excel(
        object(), frame, target, "Packing List", logs.append
    )

    waited = [line for line in logs if "vẫn đang tạo" in line]
    assert waited, logs
    assert "đã chờ 15 giây" in waited[0]
    assert target.read_bytes() == payload


def test_a_broken_cleanup_never_hides_the_file_that_was_already_downloaded(
    tmp_path, monkeypatch
):
    clock = install_fake_clock(monkeypatch, report)
    patch_automation(
        monkeypatch, report, "_report_export_url", lambda _f: "/export?Format="
    )
    payload = _workbook_bytes("BUYER INVOICE")
    frame = FetchFrame(
        [
            {
                "done": True,
                "ok": True,
                "status": 200,
                "size": len(payload),
                "prefix": "PK",
            }
        ],
        clock=clock,
        chunks=payload,
        cleanup_error=PlaywrightError("Execution context was destroyed"),
    )
    target = tmp_path / "invoice.xlsx"

    report._download_report_excel(
        object(), frame, target, "Buyer Invoice", lambda _message: None
    )

    assert frame.cleanups == 1
    assert target.read_bytes() == payload


def test_a_report_that_never_finishes_fails_instead_of_retrying_forever(
    tmp_path, monkeypatch
):
    clock = install_fake_clock(monkeypatch, report)
    patch_automation(
        monkeypatch, report, "_report_export_url", lambda _f: "/export?Format="
    )
    frame = FetchFrame([{"done": False}], clock=clock)

    with pytest.raises(RuntimeError, match=r"không hợp lệ \(HTTP unknown\)"):
        report._download_report_excel(
            object(),
            frame,
            tmp_path / "packing.xlsx",
            "Packing List",
            lambda _message: None,
        )

    assert len(frame.starts) == 1


def test_an_html_error_page_is_retried_once_then_reported_with_its_status(
    tmp_path, monkeypatch
):
    clock = install_fake_clock(monkeypatch, report)
    patch_automation(
        monkeypatch, report, "_report_export_url", lambda _f: "/export?Format="
    )
    frame = FetchFrame(
        [{"done": True, "ok": False, "status": 500, "size": 12, "prefix": "<h"}],
        clock=clock,
    )
    logs: list[str] = []

    with pytest.raises(RuntimeError, match=r"không hợp lệ \(HTTP 500\)"):
        report._download_report_excel(
            object(),
            frame,
            tmp_path / "packing.xlsx",
            "Packing List",
            logs.append,
        )

    assert len(frame.starts) == report.REPORT_DOWNLOAD_MAX_ATTEMPTS
    assert any("tải lại cùng report" in line for line in logs)


class VanishingTarget:
    """Ghi thành công nhưng file không còn trên đĩa (antivirus, sync tool)."""

    def __init__(self, path):
        self._path = path
        self.written = 0

    @property
    def parent(self):
        return self._path.parent

    def write_bytes(self, payload):
        self.written = len(payload)
        return len(payload)

    def is_file(self):
        return False

    def stat(self):  # pragma: no cover - không bao giờ tới sau is_file()
        raise AssertionError("không được đọc stat khi file đã biến mất")


def test_a_file_that_disappears_right_after_writing_is_reported_as_empty(
    tmp_path, monkeypatch
):
    clock = install_fake_clock(monkeypatch, report)
    patch_automation(
        monkeypatch, report, "_report_export_url", lambda _f: "/export?Format="
    )
    payload = _workbook_bytes("PACKING LIST")
    frame = FetchFrame(
        [
            {
                "done": True,
                "ok": True,
                "status": 200,
                "size": len(payload),
                "prefix": "PK",
            }
        ],
        clock=clock,
        chunks=payload,
    )
    target = VanishingTarget(tmp_path / "packing.xlsx")

    with pytest.raises(RuntimeError, match="bị rỗng"):
        report._download_report_excel(
            object(), frame, target, "Packing List", lambda _message: None
        )

    assert target.written == len(payload)


# --- nhận diện workbook -------------------------------------------------


def test_a_file_that_is_not_a_workbook_at_all_says_so_plainly(tmp_path):
    broken = tmp_path / "report.xlsx"
    broken.write_bytes(b"PK\x03\x04 nhung khong phai xlsx")

    with pytest.raises(RuntimeError, match="không đọc được"):
        report._report_workbook_kind(broken)


def test_identification_only_reads_the_report_header_not_the_whole_sheet(
    tmp_path,
):
    workbook = Workbook()
    sheet = workbook.active
    sheet.title = "Report"
    for index in range(1, 40):
        sheet.cell(row=index, column=1, value=f"row {index}")
    sheet.cell(row=30, column=1, value="PACKING LIST")
    path = tmp_path / "late.xlsx"
    workbook.save(path)
    workbook.close()

    # Chữ nhận diện nằm dưới dòng 20 nên không được coi là bằng chứng: hàm trả
    # rỗng và `_validate_report_kind` đi qua thay vì báo nhầm loại.
    assert report._report_workbook_kind(path) == ""
    report._validate_report_kind(path, "Buyer Invoice")


def test_a_packing_list_served_as_the_buyer_invoice_is_rejected(tmp_path):
    path = tmp_path / "swapped.xlsx"
    path.write_bytes(_workbook_bytes("PACKING LIST"))

    with pytest.raises(RuntimeError, match="trả nhầm Packing List"):
        report._validate_report_kind(path, "Buyer Invoice")


def test_a_buyer_invoice_served_as_the_packing_list_is_rejected(tmp_path):
    path = tmp_path / "swapped.xlsx"
    path.write_bytes(_workbook_bytes("COMMERCIAL INVOICE"))

    with pytest.raises(RuntimeError, match="trả nhầm Buyer Invoice"):
        report._validate_report_kind(path, "Packing List")


# --- quay lại màn Documents ---------------------------------------------


def test_returning_to_documents_does_nothing_when_the_screen_is_still_there(
    monkeypatch,
):
    clock = install_fake_clock(monkeypatch, report)
    docs = _documents_frame_node(clock)
    page = FakePage(docs, clock=clock)
    viewer = _report_frame(clock=clock)

    found_page, found_frame = report._restore_documents_screen(
        FakeContext(page), page, viewer, "https://wfx.test/docs"
    )

    assert (found_page, found_frame) == (page, docs)
    assert viewer.evaluated == []


def _restoring_context(clock, viewer, *, on_restore):
    """Context mà màn Documents chỉ quay lại sau khi automation điều hướng."""
    page = FakePage(viewer, clock=clock)
    context = FakeContext(page)

    def restore(_arg=None):
        page.frames.append(_documents_frame_node(clock))
        return None

    on_restore(restore)
    return context, page


def test_history_back_is_the_first_way_to_return_to_documents(monkeypatch):
    clock = install_fake_clock(monkeypatch, report)
    viewer = _report_frame(clock=clock)
    context, page = _restoring_context(
        clock,
        viewer,
        on_restore=lambda restore: viewer.scripts.__setitem__(
            "history.back()", restore
        ),
    )

    _found_page, found_frame = report._restore_documents_screen(
        context, page, viewer, "https://wfx.test/docs"
    )

    assert found_frame.url.endswith("SaleASNDocuments.aspx")
    assert viewer.goto_calls == []
    assert page.go_back_calls == 0


def test_a_dead_viewer_frame_falls_back_to_navigating_it_to_the_docs_url(
    monkeypatch,
):
    clock = install_fake_clock(monkeypatch, report)
    viewer = _report_frame(clock=clock)

    def explode(_arg=None):
        raise PlaywrightError("frame đã bị dispose")

    viewer.scripts["history.back()"] = explode
    page = FakePage(viewer, clock=clock)
    context = FakeContext(page)
    original_goto = viewer.goto

    def goto(url, **kwargs):
        original_goto(url, **kwargs)
        page.frames.append(_documents_frame_node(clock))

    viewer.goto = goto

    _found_page, found_frame = report._restore_documents_screen(
        context, page, viewer, "https://wfx.test/docs"
    )

    assert viewer.goto_calls == ["https://wfx.test/docs"]
    assert found_frame.url.endswith("SaleASNDocuments.aspx")


def test_when_the_frame_cannot_navigate_the_page_itself_goes_back(monkeypatch):
    clock = install_fake_clock(monkeypatch, report)

    def explode(_arg=None):
        raise PlaywrightError("frame đã bị dispose")

    viewer = _report_frame(clock=clock)
    viewer.scripts["history.back()"] = explode
    viewer.goto_error = PlaywrightError("navigation bị hủy")
    page = FakePage(viewer, clock=clock)
    context = FakeContext(page)
    original_go_back = page.go_back

    def go_back(**kwargs):
        original_go_back(**kwargs)
        page.frames.append(_documents_frame_node(clock))

    page.go_back = go_back

    _found_page, found_frame = report._restore_documents_screen(
        context, page, viewer, "https://wfx.test/docs"
    )

    assert viewer.goto_calls == ["https://wfx.test/docs"]
    assert found_frame.url.endswith("SaleASNDocuments.aspx")


def test_losing_every_way_back_reports_the_missing_documents_screen(
    monkeypatch,
):
    clock = install_fake_clock(monkeypatch, report)

    def explode(_arg=None):
        raise PlaywrightError("frame đã bị dispose")

    viewer = _report_frame(clock=clock)
    viewer.scripts["history.back()"] = explode
    viewer.goto_error = PlaywrightError("navigation bị hủy")
    page = FakePage(viewer, clock=clock)
    page.go_back_error = PlaywrightError("không còn lịch sử")
    patch_automation(
        monkeypatch, report, "DOCUMENTS_FRAME_TIMEOUT_SECONDS", 1
    )

    with pytest.raises(PlaywrightTimeoutError):
        report._restore_documents_screen(
            FakeContext(page), page, viewer, "https://wfx.test/docs"
        )

    assert page.go_back_calls == 1


# --- dọn popup ----------------------------------------------------------


def test_closing_popups_leaves_the_windows_the_user_already_had_open(
    monkeypatch,
):
    install_fake_clock(monkeypatch, report)
    kept = FakePage()
    popup = FakePage()
    already_closed = FakePage()
    already_closed.closed = True
    logs: list[str] = []

    report._close_sale_asn_document_popups(
        FakeContext(kept, popup, already_closed), {id(kept)}, logs.append
    )

    assert kept.close_calls == 0
    assert popup.close_calls == 1
    assert already_closed.close_calls == 0
    assert logs == ["[SALE ASN DOCS] Đã đóng 1 cửa sổ Docs/report."]


def test_a_popup_that_refuses_to_close_does_not_stop_the_others(monkeypatch):
    install_fake_clock(monkeypatch, report)
    stubborn = FakePage()
    stubborn.close_error = PlaywrightError("target đã đóng")
    other = FakePage()
    logs: list[str] = []

    report._close_sale_asn_document_popups(
        FakeContext(stubborn, other), set(), logs.append
    )

    assert other.close_calls == 1
    assert logs == ["[SALE ASN DOCS] Đã đóng 1 cửa sổ Docs/report."]


def test_nothing_is_logged_when_there_was_no_popup_to_close(monkeypatch):
    install_fake_clock(monkeypatch, report)
    kept = FakePage()
    logs: list[str] = []

    report._close_sale_asn_document_popups(
        FakeContext(kept), {id(kept)}, logs.append
    )

    assert logs == []
