"""Report BuyerDispatchOrder_Invoice: điền Doc No., export Excel, reload XLSX.

`wfx_panel/automation/dispatch/report.py` ở mức 22%. CLAUDE.md:

* "(GDN) Dispatch … mở report `BuyerDispatchOrder_Invoice`, điền `Doc No.`,
  chờ report load thật, export Excel Open XML rồi reload/save lại thành XLSX."
* Download phải đi qua `snapshot_downloads()` + `save_native_download()`, không
  dùng artifact tạm của Playwright.
* Report rỗng phải có mã riêng, vì đó là lỗi dữ liệu người dùng chứ không phải
  lỗi hệ thống.
"""

from __future__ import annotations

import pytest
from openpyxl import Workbook, load_workbook
from playwright.sync_api import Error as PlaywrightError
from playwright.sync_api import TimeoutError as PlaywrightTimeoutError

from tests.fakes.mini_dom import Element, MiniFrame, element
from tests.fakes.wfx_dom import install_fake_clock
from wfx_panel.automation import _common
from wfx_panel.automation.dispatch import report as dispatch_report
from wfx_panel.automation.dispatch.status import DispatchFlowError


@pytest.fixture
def clock(monkeypatch):
    return install_fake_clock(monkeypatch, dispatch_report, _common)


def _quiet():
    return lambda _line: None


# --- reload workbook ------------------------------------------------------


def _workbook(path, *, rows=((1, 2),)):
    workbook = Workbook()
    sheet = workbook.active
    for row in rows:
        sheet.append(list(row))
    workbook.save(path)
    return path


def test_a_report_workbook_is_saved_again_with_full_recalculation(tmp_path):
    source = _workbook(tmp_path / "raw.xlsx")
    target = tmp_path / "out" / "clean.xlsx"

    dispatch_report.reload_dispatch_workbook(source, target)

    assert target.is_file()
    saved = load_workbook(target)
    assert saved.calculation.fullCalcOnLoad is True
    assert saved.calculation.calcMode == "auto"


def test_a_file_that_is_not_a_workbook_is_refused(tmp_path):
    source = tmp_path / "raw.xlsx"
    source.write_bytes("không phải xlsx".encode())

    with pytest.raises(DispatchFlowError) as error:
        dispatch_report.reload_dispatch_workbook(source, tmp_path / "out.xlsx")

    assert error.value.code == "GDN_WORKBOOK_RELOAD_FAILED"
    assert error.value.errors


def test_an_empty_output_file_is_refused(tmp_path, monkeypatch):
    source = _workbook(tmp_path / "raw.xlsx")
    target = tmp_path / "out.xlsx"
    original = dispatch_report.load_workbook

    class EmptySaver:
        def __init__(self, inner):
            self._inner = inner
            self.sheetnames = inner.sheetnames
            self.worksheets = inner.worksheets
            self.calculation = inner.calculation

        def save(self, path):
            path.write_bytes(b"")

        def close(self):
            self._inner.close()

    monkeypatch.setattr(
        dispatch_report,
        "load_workbook",
        lambda *a, **kw: EmptySaver(original(*a, **kw)),
    )

    with pytest.raises(DispatchFlowError) as error:
        dispatch_report.reload_dispatch_workbook(source, target)

    assert "rỗng" in error.value.message


# --- chờ report load ------------------------------------------------------


class _ReportPage(MiniFrame):
    def __init__(self, clock, root, *, url="https://wfx.test/report"):
        super().__init__(root, url=url, clock=clock)
        self.listeners: list[tuple] = []
        self.closed = False
        self.goto_calls: list[str] = []

    def on(self, event, handler):
        self.listeners.append((event, handler))

    def remove_listener(self, event, handler):
        if (event, handler) in self.listeners:
            self.listeners.remove((event, handler))

    def close(self):
        self.closed = True

    def goto(self, url, **_kwargs):
        self.goto_calls.append(url)

    def fire_download(self, download):
        for event, handler in list(self.listeners):
            if event == "download":
                handler(download)


def _report_page(
    clock,
    *,
    image_src="export.png",
    loading=False,
    report_text="",
    export=True,
    excel=True,
    doc_no=True,
):
    children = []
    if doc_no:
        children.append(
            element("input", id="rptCustomReportViewer_ctl04_ctl03_txtValue")
        )
        children.append(element("a", id="rptCustomReportViewer_ctl04_ctl00"))
    if image_src is not None:
        children.append(
            element(
                "img",
                id="rptCustomReportViewer_ctl05_ctl04_ctl00_ButtonImg",
                attrs={"src": image_src},
            )
        )
    if loading:
        children.append(element("div", id="rptCustomReportViewer_AsyncWait"))
    if report_text:
        children.append(
            element("div", id="rptCustomReportViewer_ctl09", text=report_text)
        )
    if export:
        children.append(
            element("a", id="rptCustomReportViewer_ctl05_ctl04_ctl00_ButtonLink")
        )
    menu_children = []
    if excel:
        menu_children.append(
            element(
                "a",
                id="excel",
                attrs={"title": "Excel", "onclick": "exportReport('EXCELOPENXML')"},
            )
        )
    children.append(
        Element(
            "div",
            id="rptCustomReportViewer_ctl05_ctl04_ctl00_Menu",
            children=menu_children,
        )
    )
    return _ReportPage(clock, Element("body", children=children))


def test_a_report_with_an_enabled_export_icon_is_ready(clock):
    dispatch_report._wait_report_ready(_report_page(clock))


def test_a_report_still_showing_the_async_wait_is_not_ready(clock):
    with pytest.raises(DispatchFlowError) as error:
        dispatch_report._wait_report_ready(_report_page(clock, loading=True))

    assert error.value.code == "GDN_REPORT_NOT_READY"


def test_a_disabled_export_icon_means_the_report_is_not_ready(clock):
    with pytest.raises(DispatchFlowError):
        dispatch_report._wait_report_ready(
            _report_page(clock, image_src="export_Disabled.png")
        )


@pytest.mark.parametrize(
    "text", ["No Data", "no rows found", "Không có dữ liệu"]
)
def test_a_report_that_says_it_has_no_data_has_its_own_code(clock, text):
    page = _report_page(clock, image_src="export_Disabled.png", report_text=text)

    with pytest.raises(DispatchFlowError) as error:
        dispatch_report._wait_report_ready(page)

    assert error.value.code == "GDN_REPORT_EMPTY"


def test_a_timeout_carries_the_last_report_text(clock):
    page = _report_page(
        clock, image_src="export_Disabled.png", report_text="Đang tính toán"
    )

    with pytest.raises(DispatchFlowError) as error:
        dispatch_report._wait_report_ready(page)

    assert "Đang tính toán" in error.value.message


# --- tải Excel ------------------------------------------------------------


class _Download:
    suggested_filename = "BuyerDispatchOrder_Invoice.xlsx"


def _wire_download(monkeypatch, tmp_path, *, saves=True):
    monkeypatch.setattr(dispatch_report, "snapshot_downloads", lambda: {"before"})
    saved: list[tuple] = []

    def save(download, target, before):
        saved.append((download, target, before))
        if saves:
            target.write_bytes(b"xlsx")

    monkeypatch.setattr(dispatch_report, "save_native_download", save)
    return saved


def test_the_export_menu_is_opened_then_excel_is_clicked(
    clock, monkeypatch, tmp_path
):
    page = _report_page(clock)
    saved = _wire_download(monkeypatch, tmp_path)
    page.locator("#excel").node.on_click = lambda _n: page.fire_download(
        _Download()
    )
    target = tmp_path / "out" / "report.xlsx"
    logs: list[str] = []

    dispatch_report._download_report(page, target, logs.append)

    assert page.locator(
        "#rptCustomReportViewer_ctl05_ctl04_ctl00_ButtonLink"
    ).node.clicks == 1
    assert saved and saved[0][2] == {"before"}
    assert page.listeners == []
    assert any("export report sang Excel" in line for line in logs)


def test_a_missing_export_button_is_reported(clock, monkeypatch, tmp_path):
    page = _report_page(clock, export=False)
    _wire_download(monkeypatch, tmp_path)

    with pytest.raises(DispatchFlowError) as error:
        dispatch_report._download_report(page, tmp_path / "x.xlsx", _quiet())

    assert error.value.code == "GDN_REPORT_NOT_READY"
    assert page.listeners == []


def test_a_missing_excel_entry_is_reported(clock, monkeypatch, tmp_path):
    page = _report_page(clock, excel=False)
    _wire_download(monkeypatch, tmp_path)

    with pytest.raises(DispatchFlowError) as error:
        dispatch_report._download_report(page, tmp_path / "x.xlsx", _quiet())

    assert error.value.code == "GDN_REPORT_NOT_READY"


def test_a_download_that_never_starts_is_reported(clock, monkeypatch, tmp_path):
    page = _report_page(clock)
    _wire_download(monkeypatch, tmp_path)

    with pytest.raises(DispatchFlowError) as error:
        dispatch_report._download_report(page, tmp_path / "x.xlsx", _quiet())

    assert error.value.code == "GDN_REPORT_DOWNLOAD_FAILED"


def test_an_empty_downloaded_file_is_reported(clock, monkeypatch, tmp_path):
    page = _report_page(clock)
    _wire_download(monkeypatch, tmp_path, saves=False)
    page.locator("#excel").node.on_click = lambda _n: page.fire_download(
        _Download()
    )

    with pytest.raises(DispatchFlowError) as error:
        dispatch_report._download_report(page, tmp_path / "x.xlsx", _quiet())

    assert "rỗng" in error.value.message


# --- toàn bộ bước chuẩn bị workbook --------------------------------------


class _Context:
    def __init__(self, page):
        self._page = page
        self.new_pages = 0

    def new_page(self):
        self.new_pages += 1
        return self._page


def _wire_prepare(monkeypatch, tmp_path, *, download=True):
    monkeypatch.setattr(dispatch_report, "_wait_report_ready", lambda _page: None)

    def download_report(_page, target, _log):
        if download:
            target.parent.mkdir(parents=True, exist_ok=True)
            _workbook(target)

    monkeypatch.setattr(dispatch_report, "_download_report", download_report)


def test_the_whole_preparation_returns_the_reloaded_workbook(
    clock, monkeypatch, tmp_path
):
    page = _report_page(clock)
    context = _Context(page)
    _wire_prepare(monkeypatch, tmp_path)
    steps: list[tuple] = []

    result = dispatch_report._prepare_dispatch_workbook(
        context,
        "INV-1",
        tmp_path,
        _quiet(),
        lambda stage, _message, index, _total, state="": steps.append(
            (stage, index)
        ),
    )

    assert result.name == "BuyerDispatchOrder_Invoice.reload.xlsx"
    assert result.is_file()
    assert page.goto_calls == [dispatch_report.REPORT_URL]
    assert steps == [("report", 1), ("download", 2), ("workbook", 3)]
    assert page.closed is True


def test_a_doc_no_wfx_dropped_stops_before_the_report_runs(
    clock, monkeypatch, tmp_path
):
    page = _report_page(clock)
    page.locator(
        "#rptCustomReportViewer_ctl04_ctl03_txtValue"
    ).node.on_fill = lambda node, _value: setattr(node, "value", "")
    _wire_prepare(monkeypatch, tmp_path)

    with pytest.raises(DispatchFlowError) as error:
        dispatch_report._prepare_dispatch_workbook(
            _Context(page), "INV-1", tmp_path, _quiet()
        )

    assert error.value.code == "GDN_REPORT_NOT_READY"
    assert page.closed is True


def test_a_playwright_timeout_becomes_the_report_not_ready_code(
    clock, monkeypatch, tmp_path
):
    page = _report_page(clock)
    monkeypatch.setattr(
        dispatch_report,
        "_wait_report_ready",
        lambda _page: (_ for _ in ()).throw(PlaywrightTimeoutError("chậm")),
    )

    with pytest.raises(DispatchFlowError) as error:
        dispatch_report._prepare_dispatch_workbook(
            _Context(page), "INV-1", tmp_path, _quiet()
        )

    assert error.value.code == "GDN_REPORT_NOT_READY"
    assert error.value.errors


def test_a_playwright_failure_becomes_the_download_failed_code(
    clock, monkeypatch, tmp_path
):
    page = _report_page(clock)
    monkeypatch.setattr(
        dispatch_report,
        "_wait_report_ready",
        lambda _page: (_ for _ in ()).throw(PlaywrightError("tab rơi")),
    )

    with pytest.raises(DispatchFlowError) as error:
        dispatch_report._prepare_dispatch_workbook(
            _Context(page), "INV-1", tmp_path, _quiet()
        )

    assert error.value.code == "GDN_REPORT_DOWNLOAD_FAILED"


def test_a_dispatch_error_is_passed_through_unchanged(
    clock, monkeypatch, tmp_path
):
    page = _report_page(clock)
    monkeypatch.setattr(
        dispatch_report,
        "_wait_report_ready",
        lambda _page: (_ for _ in ()).throw(
            DispatchFlowError("GDN_REPORT_EMPTY", "trống")
        ),
    )

    with pytest.raises(DispatchFlowError) as error:
        dispatch_report._prepare_dispatch_workbook(
            _Context(page), "INV-1", tmp_path, _quiet()
        )

    assert error.value.code == "GDN_REPORT_EMPTY"


def test_the_report_tab_is_always_closed(clock, monkeypatch, tmp_path):
    page = _report_page(clock)
    monkeypatch.setattr(
        dispatch_report,
        "_wait_report_ready",
        lambda _page: (_ for _ in ()).throw(PlaywrightError("tab rơi")),
    )

    with pytest.raises(DispatchFlowError):
        dispatch_report._prepare_dispatch_workbook(
            _Context(page), "INV-1", tmp_path, _quiet()
        )

    assert page.closed is True
