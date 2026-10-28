from datetime import date
from pathlib import Path
from types import SimpleNamespace

import pytest
from openpyxl import Workbook, load_workbook

from wfx_panel.app.dialogs import FileDialogController
from wfx_panel.automation import expense_invoice as automation
from wfx_panel.controllers import finance
from wfx_panel.workbooks.expense_invoice import (
    COLUMNS,
    OPTION_KEYS,
    SHEET,
    TEMPLATE,
    ExpenseWorkbookError,
    load_seed_options,
    read_expense_workbook,
    save_template,
)


def input_file(tmp_path, *, second=False, edit=None):
    workbook = Workbook()
    sheet = workbook.active
    sheet.title = SHEET
    sheet.append(COLUMNS)
    sheet.append(
        [
            "Supplier A",
            "00012",
            date(2026, 10, 28),
            date(2026, 10, 28),
            "USD",
            "TT Before Shipment",
            date(2026, 11, 28),
            "6421901",
            2,
            3.125,
            "Ghi chú",
            "WFX",
            "Vietnam",
        ]
    )
    if second:
        sheet.append([None] * 7 + ["6421902", 1, 5, None, None, None])
    if edit:
        edit(sheet)
    path = tmp_path / "input.xlsx"
    workbook.save(path)
    workbook.close()
    return path


def test_multiple_ledgers_inherit_header_and_keep_invoice_identifier(tmp_path):
    data = read_expense_workbook(input_file(tmp_path, second=True))
    assert data["header"]["invoice_no"] == "00012"
    assert data["header"]["invoice_date"] == "2026-10-28"
    assert len(data["lines"]) == 2
    assert data["total"] == "11.250"


@pytest.mark.parametrize(
    "cell,value",
    [
        ("I2", 0),
        ("I2", -1),
        ("I2", 1.001),
        ("J2", -1),
        ("J2", "NaN"),
        ("J2", 1.00001),
        ("J2", "=1+1"),
        ("C2", "31/02/2026"),
        ("A2", None),
        ("B3", "different invoice"),
        ("L3", "different brand"),
    ],
)
def test_reject_invalid_or_conflicting_data_before_browser(tmp_path, cell, value):
    path = input_file(
        tmp_path, second=True, edit=lambda s: setattr(s[cell], "value", value)
    )
    with pytest.raises(ExpenseWorkbookError):
        read_expense_workbook(path)


def test_zero_rate_allowed_and_text_dates_are_normalized(tmp_path):
    def edit(sheet):
        sheet["J2"] = 0
        sheet["C2"] = "28/10/2026"

    data = read_expense_workbook(input_file(tmp_path, edit=edit))
    assert data["total"] == "0"
    assert data["header"]["invoice_date"] == "2026-10-28"


def test_template_dropdowns_use_hidden_named_ranges_and_keep_13_columns(tmp_path):
    assert TEMPLATE.is_file()
    options = {
        key: ["Choice 1", "Choice 2"]
        for key in ("supplier", "currency", "payment_term", "ledger", "brand", "team")
    }
    target = tmp_path / "template.xlsx"
    save_template(target, options)
    wb = load_workbook(target)
    try:
        assert tuple(c.value for c in wb[SHEET][1]) == COLUMNS
        assert wb["_Lists"].sheet_state == "veryHidden"
        dropdowns = [
            v for v in wb[SHEET].data_validations.dataValidation if v.type == "list"
        ]
        assert len(dropdowns) == 6
        for validation in dropdowns:
            name = validation.formula1.removeprefix("=")
            tab, area = next(wb.defined_names[name].destinations)
            assert [c[0].value for c in wb[tab][area]] == ["Choice 1", "Choice 2"]
        assert {v.type for v in wb[SHEET].data_validations.dataValidation} >= {
            "date",
            "decimal",
        }
    finally:
        wb.close()
    with pytest.raises(ExpenseWorkbookError, match="chưa có dòng"):
        read_expense_workbook(target)


def controller(tmp_path):
    panel = SimpleNamespace(
        _base_dir=tmp_path,
        _account=lambda: {"user_id": "tester"},
        _current_division="A",
        _log=lambda m: None,
    )
    panel._run = lambda method, action, *args: action()
    return finance.FinanceController(panel)


def test_review_snapshots_data_and_is_single_use(tmp_path, monkeypatch):
    ctrl = controller(tmp_path)
    path = input_file(tmp_path)
    review = ctrl.review_expense_import(str(path))
    path.write_bytes(b"file changed after review")
    received = []
    monkeypatch.setattr(
        finance,
        "fill_expense_invoice",
        lambda data, log: received.append(data) or {"ok": True},
    )
    assert ctrl.fill_expense_import(review["review_token"])["ok"]
    assert received[0]["header"]["invoice_no"] == "00012"
    assert (
        ctrl.fill_expense_import(review["review_token"])["code"]
        == "EXPENSE_REVIEW_EXPIRED"
    )
    assert len(received) == 1


def test_division_change_expires_review_without_touching_browser(tmp_path, monkeypatch):
    ctrl = controller(tmp_path)
    review = ctrl.review_expense_import(str(input_file(tmp_path)))
    ctrl._panel._current_division = "B"
    monkeypatch.setattr(
        finance,
        "fill_expense_invoice",
        lambda *a: pytest.fail("browser must not be touched"),
    )
    assert not ctrl.fill_expense_import(review["review_token"])["ok"]


def test_safe_session_recovery_can_reuse_consumed_snapshot_inside_same_run(
    tmp_path, monkeypatch
):
    ctrl = controller(tmp_path)
    review = ctrl.review_expense_import(str(input_file(tmp_path)))
    results = iter([{"ok": False, "code": "NOT_LOGGED_IN"}, {"ok": True}])
    monkeypatch.setattr(finance, "fill_expense_invoice", lambda *args: next(results))

    def run_with_recovery(method, action, *args):
        assert action()["code"] == "NOT_LOGGED_IN"
        return action()

    ctrl._panel._run = run_with_recovery
    assert ctrl.fill_expense_import(review["review_token"])["ok"]
    assert not ctrl.expense_reviews


def test_ledger_exact_code_only_and_ambiguous_options_rejected():
    options = [{"value": "x", "label": "6421901(6421901-WFX SUBSCRIPTION)"}]
    assert automation._option(options, "6421901", ledger=True)["value"] == "x"
    with pytest.raises(automation.ExpenseInputError):
        automation._option(options, "64219", ledger=True)
    with pytest.raises(automation.ExpenseInputError):
        automation._option(
            options + [{"value": "y", "label": options[0]["label"]}],
            "6421901",
            ledger=True,
        )


HTML = """<!doctype html><h1>Expense Invoice - New</h1>
<button onclick="window.saved=true">Save</button><button onclick="window.confirmed=true">Confirm</button>
<select id=ddlInvoiceType><option value=GeneralExpense>General Expense</option></select>
<select id=ddlSubType><option value=other>ALL OTHER EXPENSES</option></select>
<select id=ddlSupplier><option value=''>[Select]</option><option value=s>Supplier A</option></select>
<select id=ddlInvoiceCurrency><option value=USD>USD</option></select>
<select id=ddlPaymentTerms><option value=tt>TT Before Shipment</option></select>
<select id=ddlBrand><option value=''>[Select]</option><option value=wfx>WFX</option></select>
<select id=ddlTeam><option value=''>[Select]</option><option value=VN>Vietnam</option></select>
<input id=txtInvoiceNo><input id=dtInvoiceDate><input id=dtGLDate><input id=dtDueDate><textarea id=txtComments></textarea>
<a href='#' onclick='addRow();return false'>Add</a>
<table id=gridExpenseInvoiceExpenseDetail_tblGridHeader><tr><td>Ledger Qty Rate</td></tr></table>
<table id=gridExpenseInvoiceExpenseDetail_tblGridContent></table>
<script>
// Handler dropdown của WFX (HdlrWFXDropDown.ashx) chỉ trả dữ liệu, không đổi form.
const GUID='guid';
window.gInvoiceType='GeneralExpense';
function GetBindDDLData(guid, type, page, params){
 window.handlerCalls=(window.handlerCalls||[]).concat([type+'|'+params]);
 return {
  Supplier:[{SupplierCompanyName:'Supplier A',CurrencyCode:'USD'},{SupplierCompanyName:'Supplier B',CurrencyCode:'VND'},{SupplierCompanyName:'Supplier C',CurrencyCode:'USD'}],
  Ledger:[{LedgerAcName:'6421901(6421901-WFX SUBSCRIPTION)',Active:'1'},{LedgerAcName:'6421902(6421902-HR SUBSCRIPTION)',Active:'1'},{LedgerAcName:'9999(OLD)',Active:'0'}],
  PaymentTerms:[{PaymentTerm:'TT Before Shipment'},{PaymentTerm:'30 Days'}],
 }[type]||[];
}
function addRow(){
 const row=document.querySelector('#gridExpenseInvoiceExpenseDetail_tblGridContent').insertRow();
 for (const [id, editor, initial] of [['ColLedgerAcName','ddlLedgerAcName',''],['ColExpenseQty','txtExpenseQty','0'],['ColExpenseRate','txtExpenseRate','0']]){
  const cell=row.insertCell();cell.id=id;cell.textContent=initial;
  cell.ondblclick=()=>{
   if(cell.querySelector('input,select'))return;
   const old=cell.textContent;cell.textContent='';
   const input=document.createElement(editor==='ddlLedgerAcName'?'select':'input');input.id=editor;
   if(editor==='ddlLedgerAcName'){
    input.innerHTML='<option value="">[Select]</option><option value="ledger1">6421901(6421901-WFX SUBSCRIPTION)</option><option value="ledger2">6421902(6421902-HR SUBSCRIPTION)</option>';
    input.value=old.startsWith('6421901')?'ledger1':old.startsWith('6421902')?'ledger2':'';
   }else input.value=old;
   cell.append(input);input.focus();
   input.onblur=()=>{cell.textContent=editor==='ddlLedgerAcName'?input.selectedOptions[0].textContent:input.value;};
  };
 }
}
addRow();
</script>"""


@pytest.fixture
def expense_page(monkeypatch):
    from playwright.sync_api import sync_playwright

    chrome = Path("C:/Program Files/Google/Chrome/Application/chrome.exe")
    with sync_playwright() as playwright:
        browser = (
            playwright.chromium.launch(executable_path=str(chrome), headless=True)
            if chrome.exists()
            else playwright.chromium.launch(headless=True)
        )
        context = browser.new_context()
        page = context.new_page()
        page.route(
            "**/*", lambda route: route.fulfill(body=HTML, content_type="text/html")
        )
        page.goto("http://wfx.test/WFXExpenseInvoice.aspx?InvoiceType=Expense")
        monkeypatch.setattr(
            automation,
            "sync_playwright",
            lambda: SimpleNamespace(start=lambda: SimpleNamespace(stop=lambda: None)),
        )
        monkeypatch.setattr(
            automation, "_active_wfx_page", lambda *args: (browser, page)
        )
        yield page
        browser.close()


def test_fill_multiple_rows_in_browser_and_never_save_or_confirm(
    tmp_path, expense_page
):
    data = read_expense_workbook(input_file(tmp_path, second=True))
    result = automation.fill_expense_invoice(data, log=lambda m: None)
    assert result["code"] == "EXPENSE_FORM_COMPLETED", result
    assert expense_page.locator("#txtInvoiceNo").input_value() == "00012"
    assert expense_page.locator("#ddlSupplier").input_value() == "s"
    assert expense_page.locator("#dtDueDate").input_value() == "28 Nov 2026"
    assert automation._rows(expense_page).count() == 2
    assert not expense_page.evaluate("Boolean(window.saved || window.confirmed)")
    # Một lần chạy tiếp không được ghi đè draft vừa tạo.
    again = automation.fill_expense_invoice(data, log=lambda m: None)
    assert not again["ok"] and not again["partially_filled"]


def test_fill_can_use_same_supplier_already_selected_on_otherwise_empty_new(
    tmp_path, expense_page
):
    expense_page.locator("#ddlSupplier").select_option("s")
    result = automation.fill_expense_invoice(
        read_expense_workbook(input_file(tmp_path)), log=lambda m: None
    )
    assert result["ok"], result


def test_read_options_lists_everything_without_supplier_or_form_changes(expense_page):
    expense_page.evaluate("""() => {
        const currency=document.querySelector('#ddlInvoiceCurrency');
        currency.innerHTML='<option value="">[Select]</option>';
        // WFX bật hộp "Please select Supplier first" khi mousedown lúc này.
        currency.onmousedown=()=>{window.supplierPrompt=true;};
        document.querySelector('#gridExpenseInvoiceExpenseDetail_tblGridContent').innerHTML='';
    }""")
    options = automation.read_expense_options(expense_page)
    assert options == {
        "supplier": ["Supplier A", "Supplier B", "Supplier C"],
        "currency": ["USD", "VND"],
        "payment_term": ["TT Before Shipment", "30 Days"],
        "ledger": [
            "6421901(6421901-WFX SUBSCRIPTION)",
            "6421902(6421902-HR SUBSCRIPTION)",
        ],
        "brand": ["WFX"],
        "team": ["Vietnam"],
    }
    assert "Ledger|InvoiceSubType|other~" in expense_page.evaluate("window.handlerCalls")
    assert not expense_page.evaluate("Boolean(window.supplierPrompt)")
    assert expense_page.locator("#ddlSupplier").input_value() == ""
    assert not expense_page.locator("#ddlLedgerAcName").count()
    assert not expense_page.evaluate("Boolean(window.saved || window.confirmed)")


def test_read_options_refuses_incomplete_lists(expense_page):
    expense_page.evaluate("() => { window.GetBindDDLData = () => []; }")
    with pytest.raises(automation.ExpenseInputError, match="supplier"):
        automation.read_expense_options(expense_page)


def test_successful_fill_refreshes_dropdown_cache_for_this_account(
    tmp_path, expense_page
):
    import json

    ctrl = controller(tmp_path)
    review = ctrl.review_expense_import(str(input_file(tmp_path)))
    result = ctrl.fill_expense_import(review["review_token"])
    assert result["code"] == "EXPENSE_FORM_COMPLETED", result
    assert "options" not in result
    cached = json.loads((tmp_path / "expense-options.json").read_text(encoding="utf-8"))
    assert cached["context"] == ["tester", "A"]
    assert cached["options"]["supplier"] == ["Supplier A", "Supplier B", "Supplier C"]
    template = ctrl.expense_template_options()
    assert template["code"] == "EXPENSE_OPTIONS_CACHED"
    ctrl._panel._current_division = "B"
    assert ctrl.expense_template_options()["code"] == "EXPENSE_OPTIONS_SEED"


def test_options_refresh_failure_never_fails_the_fill(tmp_path, expense_page):
    expense_page.evaluate("() => { window.GetBindDDLData = () => []; }")
    result = automation.fill_expense_invoice(
        read_expense_workbook(input_file(tmp_path)), log=lambda m: None
    )
    assert result["code"] == "EXPENSE_FORM_COMPLETED", result
    assert result["options"] is None


def test_seed_has_full_dropdowns_for_every_list_column():
    seed = load_seed_options()
    for key in OPTION_KEYS:
        assert seed.get(key), key


def test_template_options_never_touch_browser(tmp_path, monkeypatch):
    ctrl = controller(tmp_path)
    ctrl._panel._run = lambda *a: pytest.fail("template must not run automation")
    result = ctrl.expense_template_options()
    assert result["code"] == "EXPENSE_OPTIONS_SEED"
    assert result["options"] == load_seed_options()
    monkeypatch.setattr(finance, "load_seed_options", lambda: {})
    result = ctrl.expense_template_options()
    assert result["ok"] and result["code"] == "EXPENSE_OPTIONS_UNAVAILABLE"


def test_template_saves_immediately_without_choosing_supplier(tmp_path, monkeypatch):
    target = tmp_path / "selected.xlsx"
    ctrl = controller(tmp_path)
    app = SimpleNamespace(
        window=SimpleNamespace(create_file_dialog=lambda *a, **k: [str(target)]),
        api=SimpleNamespace(_finance=ctrl, _log=lambda m: None),
    )
    dialogs = FileDialogController(app)
    monkeypatch.setattr(dialogs, "_handle_downloaded_excel", lambda *args: True)
    result = dialogs.expense_file_dialog(save=True)
    assert result["code"] == "EXPENSE_TEMPLATE_SAVED", result
    wb = load_workbook(target)
    try:
        assert wb[SHEET]["A2"].value is None
        dropdowns = [
            v for v in wb[SHEET].data_validations.dataValidation if v.type == "list"
        ]
        assert len(dropdowns) == 6
        assert {v.errorStyle for v in dropdowns} == {"warning"}
    finally:
        wb.close()


def test_template_without_any_list_still_saves(tmp_path):
    target = tmp_path / "plain.xlsx"
    save_template(target, {})
    wb = load_workbook(target)
    try:
        assert tuple(c.value for c in wb[SHEET][1]) == COLUMNS
        assert not [
            v for v in wb[SHEET].data_validations.dataValidation if v.type == "list"
        ]
    finally:
        wb.close()


def test_ledger_editor_can_commit_and_detach_on_change(tmp_path, expense_page):
    expense_page.evaluate("""() => {
        document.addEventListener('change', e => {
            if (e.target.id === 'ddlLedgerAcName') {
                const field = e.target, label=field.selectedOptions[0].textContent;
                field.onblur=null;
                field.parentElement.textContent=label;
            }
        });
    }""")
    data = read_expense_workbook(input_file(tmp_path))
    result = automation.fill_expense_invoice(data, log=lambda m: None)
    assert result["code"] == "EXPENSE_FORM_COMPLETED", result


def test_supplier_reset_by_later_callback_must_not_be_reported_as_success(
    tmp_path, expense_page
):
    expense_page.evaluate("""() => {
        document.querySelector('#txtInvoiceNo').addEventListener('change', () => {
            document.querySelector('#ddlSupplier').value='';
        });
    }""")
    data = read_expense_workbook(input_file(tmp_path))
    result = automation.fill_expense_invoice(data, log=lambda m: None)
    assert not result["ok"] and result["partially_filled"]


def test_optional_blank_flex_fields_clear_supplier_defaults(tmp_path, expense_page):
    expense_page.evaluate("""() => {
        document.querySelector('#ddlSupplier').addEventListener('change', () => {
            document.querySelector('#ddlBrand').value='wfx';
            document.querySelector('#ddlTeam').value='VN';
        });
    }""")

    def edit(sheet):
        sheet["L2"] = None
        sheet["M2"] = None

    result = automation.fill_expense_invoice(
        read_expense_workbook(input_file(tmp_path, edit=edit)), log=lambda m: None
    )
    assert result["ok"], result
    assert expense_page.locator("#ddlBrand").input_value() == ""
    assert expense_page.locator("#ddlTeam").input_value() == ""
