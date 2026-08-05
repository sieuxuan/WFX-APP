from pathlib import Path
from zipfile import ZipFile

import pytest
from openpyxl import load_workbook

from wfx_panel.oc_workbook import write_oc_input_template
from wfx_panel.sale_asn_workbook import write_sale_asn_template
from wfx_panel.style_workbook import write_style_template


@pytest.mark.parametrize(
    ("file_name", "writer"),
    (
        ("style.xlsx", write_style_template),
        ("oc.xlsx", write_oc_input_template),
        ("sale-asn.xlsx", write_sale_asn_template),
    ),
)
def test_downloadable_templates_have_excel_safe_filters_and_tables(
    tmp_path: Path,
    file_name,
    writer,
):
    target = writer(tmp_path / file_name)

    with ZipFile(target) as archive:
        assert archive.testzip() is None

    workbook = load_workbook(target)
    try:
        for sheet in workbook.worksheets:
            tables = list(sheet.tables.values())
            if tables:
                # Một worksheet không được có AutoFilter riêng chồng lên
                # AutoFilter nằm sẵn trong Excel Table.
                assert sheet.auto_filter.ref is None
            for table in tables:
                headers = [column.name for column in table.tableColumns]
                assert headers
                assert len(headers) == len(set(headers))
                assert all(str(header).strip() for header in headers)
                assert table.autoFilter is not None
                assert table.autoFilter.ref == table.ref
    finally:
        workbook.close()
