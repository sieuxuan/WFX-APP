"""Kiểm tra file Upload OC trước khi mở WFX: an toàn ZIP, giá trị và ngày.

CLAUDE.md: trước khi mở WFX phải kiểm extension/ZIP an toàn, schema/header, ô
lỗi công thức, trường bắt buộc, ngày và Selling Price. Lỗi ở đây là lỗi người
dùng nên không gửi telemetry — càng phải báo chính xác dòng nào sai.
"""

from __future__ import annotations

from datetime import date, datetime
from decimal import Decimal
from io import BytesIO
from zipfile import ZIP_DEFLATED, ZipFile

import pytest
from openpyxl import Workbook

from wfx_panel.workbooks.oc import OCWorkbookError
from wfx_panel.workbooks.oc import values as oc_values
from wfx_panel.workbooks.oc.schema import EDI_HEADERS, INPUT_HEADERS


def _xlsx(tmp_path, name="OC.xlsx", *, sheets=None):
    workbook = Workbook()
    first = True
    for sheet_name, rows in (sheets or {"Sheet1": [list(EDI_HEADERS)]}).items():
        sheet = workbook.active if first else workbook.create_sheet()
        sheet.title = sheet_name
        first = False
        for row in rows:
            sheet.append(list(row))
    path = tmp_path / name
    workbook.save(path)
    workbook.close()
    return path


# --- an toàn archive ----------------------------------------------------


def test_only_xlsx_is_accepted(tmp_path):
    path = tmp_path / "OC.xls"
    path.write_bytes(b"excel")

    with pytest.raises(OCWorkbookError) as error:
        oc_values._validate_xlsx_archive(path)

    assert error.value.code == "OC_FILE_TYPE_UNSUPPORTED"


def test_a_file_the_user_already_moved_is_reported(tmp_path):
    with pytest.raises(OCWorkbookError) as error:
        oc_values._validate_xlsx_archive(tmp_path / "khong-ton-tai.xlsx")

    assert error.value.code == "OC_FILE_NOT_FOUND"


def test_a_file_above_the_size_cap_is_refused(tmp_path, monkeypatch):
    path = _xlsx(tmp_path)
    monkeypatch.setattr(oc_values, "MAX_XLSX_BYTES", 10)

    with pytest.raises(OCWorkbookError) as error:
        oc_values._validate_xlsx_archive(path)

    assert error.value.code == "OC_FILE_TOO_LARGE"


def test_a_file_that_is_not_a_zip_at_all_is_reported_as_invalid(tmp_path):
    path = tmp_path / "OC.xlsx"
    path.write_bytes(b"khong phai zip")

    with pytest.raises(OCWorkbookError) as error:
        oc_values._validate_xlsx_archive(path)

    assert error.value.code == "OC_FILE_INVALID"


def test_an_archive_with_too_many_entries_is_refused(tmp_path, monkeypatch):
    path = _xlsx(tmp_path)
    monkeypatch.setattr(oc_values, "MAX_ARCHIVE_ENTRIES", 1)

    with pytest.raises(OCWorkbookError) as error:
        oc_values._validate_xlsx_archive(path)

    assert error.value.code == "OC_FILE_UNSAFE"


def _archive(tmp_path, entries, name="OC.xlsx"):
    path = tmp_path / name
    with ZipFile(path, "w", ZIP_DEFLATED) as archive:
        for entry_name, payload in entries:
            archive.writestr(entry_name, payload)
    return path


@pytest.mark.parametrize(
    "entry_name",
    ["../escape.xml", "/absolute.xml", "xl/../../escape.xml"],
)
def test_an_archive_carrying_a_path_traversal_entry_is_refused(
    tmp_path, entry_name
):
    path = _archive(
        tmp_path,
        [
            ("[Content_Types].xml", "<Types/>"),
            ("xl/workbook.xml", "<workbook/>"),
            (entry_name, "x"),
        ],
    )

    with pytest.raises(OCWorkbookError) as error:
        oc_values._validate_xlsx_archive(path)

    assert error.value.code == "OC_FILE_UNSAFE"


def test_a_zip_bomb_is_refused_before_it_is_expanded(tmp_path, monkeypatch):
    path = _archive(
        tmp_path,
        [
            ("[Content_Types].xml", "<Types/>"),
            ("xl/workbook.xml", "<workbook/>"),
            ("xl/big.xml", "x" * 5_000),
        ],
    )
    monkeypatch.setattr(oc_values, "MAX_UNCOMPRESSED_BYTES", 1_000)

    with pytest.raises(OCWorkbookError) as error:
        oc_values._validate_xlsx_archive(path)

    assert error.value.code == "OC_FILE_UNSAFE"


def test_a_zip_that_is_not_an_excel_workbook_is_refused(tmp_path):
    path = _archive(tmp_path, [("readme.txt", "chi la file zip")])

    with pytest.raises(OCWorkbookError) as error:
        oc_values._validate_xlsx_archive(path)

    assert error.value.code == "OC_FILE_INVALID"


def test_a_real_workbook_passes_every_archive_check(tmp_path):
    oc_values._validate_xlsx_archive(_xlsx(tmp_path))


# --- không được có công thức -------------------------------------------


def test_a_workbook_that_cannot_be_opened_is_reported(tmp_path):
    path = tmp_path / "OC.xlsx"
    path.write_bytes(b"PK\x03\x04 nhung khong phai xlsx")

    with pytest.raises(OCWorkbookError) as error:
        oc_values._ensure_input_values_only(path, "new")

    assert error.value.code == "OC_FILE_INVALID"


def test_a_revise_file_is_checked_on_sheet1(tmp_path):
    workbook = Workbook()
    sheet = workbook.active
    sheet.title = "Sheet1"
    sheet.append(list(EDI_HEADERS))
    sheet["A2"] = "=1+1"
    path = tmp_path / "OC.xlsx"
    workbook.save(path)
    workbook.close()

    with pytest.raises(OCWorkbookError) as error:
        oc_values._ensure_input_values_only(path, "revise")

    assert error.value.code == "OC_FILE_FORMULA_ERROR"
    assert "A2" in error.value.errors[0]


def test_a_new_file_is_checked_on_the_input_sheet(tmp_path):
    workbook = Workbook()
    sheet = workbook.active
    sheet.title = oc_values.INPUT_SHEET_NAME
    sheet.append(list(INPUT_HEADERS))
    sheet["A2"] = "=1+1"
    path = tmp_path / "OC.xlsx"
    workbook.save(path)
    workbook.close()

    with pytest.raises(OCWorkbookError) as error:
        oc_values._ensure_input_values_only(path, "new")

    assert error.value.code == "OC_FILE_FORMULA_ERROR"


def test_more_than_twelve_formula_cells_are_summarised(tmp_path):
    workbook = Workbook()
    sheet = workbook.active
    sheet.title = "Sheet1"
    sheet.append(list(EDI_HEADERS))
    for index in range(2, 20):
        sheet.cell(row=index, column=1, value="=1+1")
    path = tmp_path / "OC.xlsx"
    workbook.save(path)
    workbook.close()

    with pytest.raises(OCWorkbookError) as error:
        oc_values._ensure_input_values_only(path, "revise")

    assert error.value.errors[0].rstrip().endswith("…")


def test_a_workbook_without_the_expected_sheet_is_simply_skipped(tmp_path):
    path = _xlsx(tmp_path, sheets={"KHAC": [["a"]]})

    oc_values._ensure_input_values_only(path, "revise")


def test_an_old_form_workbook_is_checked_on_its_form_sheet(tmp_path):
    workbook = Workbook()
    sheet = workbook.active
    sheet.title = "FORM"
    sheet["A1"] = "=1+1"
    path = tmp_path / "OC.xlsx"
    workbook.save(path)
    workbook.close()

    with pytest.raises(OCWorkbookError) as error:
        oc_values._ensure_input_values_only(path, "new")

    assert error.value.code == "OC_FILE_FORMULA_ERROR"


# --- số -----------------------------------------------------------------


@pytest.mark.parametrize(
    ("value", "expected"),
    [("1,250.50", Decimal("1250.50")), (12, Decimal("12")), (" 3 ", Decimal("3"))],
)
def test_numbers_are_read_the_way_excel_shows_them(value, expected):
    assert oc_values._decimal(value, "Units", 2) == expected


@pytest.mark.parametrize("value", ["abc", None, "", "1.2.3"])
def test_anything_that_is_not_a_number_names_its_row_and_column(value):
    with pytest.raises(OCWorkbookError) as error:
        oc_values._decimal(value, "Units", 7)

    assert error.value.code == "OC_FILE_VALIDATION_FAILED"
    assert "Dòng 7: Units phải là số." in error.value.errors


@pytest.mark.parametrize("value", [True, False])
def test_an_excel_boolean_cell_is_a_file_error_not_a_crash(value):
    """TRUE/FALSE trong ô Units phải ra lỗi file, không thoát thành lỗi thô.

    Thoát ra ngoài thì `_run` quy về PANEL_ERROR — mã nằm ngoài
    NON_REPORTABLE_FAILURES — nên telemetry sẽ gửi một lỗi hệ thống không có
    thật ra webhook production cho một lỗi nhập liệu của người dùng.
    """
    from wfx_panel.run_policy import NON_REPORTABLE_FAILURES

    with pytest.raises(OCWorkbookError) as error:
        oc_values._decimal(value, "Units", 7)

    assert error.value.code in NON_REPORTABLE_FAILURES
    assert "Dòng 7: Units phải là số." in error.value.errors


@pytest.mark.parametrize("value", ["NaN", "Infinity", "-Infinity"])
def test_a_number_excel_cannot_really_hold_is_refused(value):
    with pytest.raises(OCWorkbookError) as error:
        oc_values._decimal(value, "Units", 7)

    assert "hữu hạn" in error.value.errors[0]


# --- ngày ---------------------------------------------------------------


@pytest.mark.parametrize(
    "value",
    [
        date(2026, 3, 1),
        datetime(2026, 3, 1, 10, 30),
        "01-03-2026",
        "01/03/2026",
        "2026-03-01",
    ],
)
def test_every_date_format_the_form_accepts_reads_as_the_same_day(value):
    assert oc_values._date_value(value, "Buyer Order Date", 2) == date(2026, 3, 1)


def test_an_excel_serial_number_is_converted():
    assert oc_values._date_value(46082, "Buyer Order Date", 2) == date(2026, 3, 1)


def test_a_serial_number_excel_cannot_convert_falls_through_to_text():
    with pytest.raises(OCWorkbookError) as error:
        oc_values._date_value(-10_000_000, "Buyer Order Date", 4)

    assert "phải là ngày hợp lệ" in error.value.errors[0]


@pytest.mark.parametrize("value", ["khong phai ngay", "", None, True])
def test_anything_that_is_not_a_date_names_its_row(value):
    with pytest.raises(OCWorkbookError) as error:
        oc_values._date_value(value, "Buyer Delivery Date", 5)

    assert "Dòng 5: Buyer Delivery Date phải là ngày hợp lệ." in error.value.errors


# --- văn bản an toàn ----------------------------------------------------


@pytest.mark.parametrize("value", ["=cmd()", "+1", "@SUM(A1)"])
def test_a_cell_that_starts_like_a_formula_is_refused(value):
    with pytest.raises(OCWorkbookError) as error:
        oc_values._safe_text(value, "Buyer", 3)

    assert "không được bắt đầu bằng công thức" in error.value.errors[0]


@pytest.mark.parametrize(
    "value", ["#REF!", "#VALUE!", "#N/A", "#NAME?", "#DIV/0!", "#NUM!", "#NULL!"]
)
def test_an_excel_error_cell_is_reported_as_a_formula_error(value):
    with pytest.raises(OCWorkbookError) as error:
        oc_values._safe_text(value, "Buyer", 3)

    assert error.value.code == "OC_FILE_FORMULA_ERROR"


def test_a_hash_that_is_not_an_excel_error_is_ordinary_text():
    assert oc_values._safe_text("#1 Jacket", "Style", 3) == "#1 Jacket"


def test_whitespace_is_collapsed_so_wfx_matching_stays_stable():
    assert oc_values._safe_text("  J.LINDEBERG   AB ", "Buyer", 3) == (
        "J.LINDEBERG AB"
    )


# --- dòng Units = 0 -----------------------------------------------------


@pytest.mark.parametrize("value", [0, "0", "0.00", Decimal("0")])
def test_a_zero_quantity_row_is_recognised(value):
    assert oc_values._is_zero_quantity(value) is True


@pytest.mark.parametrize("value", [None, "", True, False, "abc", 1, "-1"])
def test_anything_else_is_not_a_zero_quantity_row(value):
    assert oc_values._is_zero_quantity(value) is False


# --- danh sách chuẩn ----------------------------------------------------


OPTIONS = ("Confirmed", "Forecast", "SMS")


def test_an_option_is_matched_ignoring_case_and_returned_canonical():
    assert (
        oc_values._canonical_option("  forecast ", "Order Type", 2, OPTIONS)
        == "Forecast"
    )


def test_an_empty_cell_falls_back_to_the_default_when_there_is_one():
    assert (
        oc_values._canonical_option("", "Zone", 2, ("FOB", "EXW"), default="FOB")
        == "FOB"
    )


def test_an_empty_cell_without_a_default_is_an_error():
    with pytest.raises(OCWorkbookError) as error:
        oc_values._canonical_option("", "Order Type", 9, OPTIONS)

    assert "Dòng 9" in error.value.errors[0]
    assert "[trống]" in error.value.errors[0]


def test_a_value_outside_the_list_names_what_the_user_typed():
    with pytest.raises(OCWorkbookError) as error:
        oc_values._canonical_option("La Lam", "Order Type", 9, OPTIONS)

    assert "La Lam" in error.value.errors[0]


# --- thứ tự ngày giao ---------------------------------------------------


def test_the_required_date_order_passes_when_it_is_respected():
    assert (
        oc_values._validate_delivery_dates(
            date(2026, 1, 1), date(2026, 2, 1), date(2026, 3, 1), 2
        )
        == []
    )


@pytest.mark.parametrize(
    ("order", "eta", "delivery"),
    [
        (date(2026, 2, 1), date(2026, 1, 1), date(2026, 3, 1)),
        (date(2026, 1, 1), date(2026, 3, 1), date(2026, 2, 1)),
        (date(2026, 1, 1), date(2026, 1, 1), date(2026, 3, 1)),
    ],
)
def test_a_broken_date_order_is_reported_with_its_row(order, eta, delivery):
    errors = oc_values._validate_delivery_dates(order, eta, delivery, 4)

    assert errors and "Dòng 4" in errors[0]


def test_the_oc_delivery_date_must_equal_the_buyer_delivery_date():
    errors = oc_values._validate_delivery_dates(
        date(2026, 1, 1),
        date(2026, 2, 1),
        date(2026, 3, 1),
        4,
        oc_delivery_date=date(2026, 3, 2),
    )

    assert errors == [
        "Dòng 4: Buyer Delivery Date phải bằng OC Delivery Date."
    ]


# --- header -------------------------------------------------------------


def test_matching_headers_pass_even_with_stray_spaces_and_dots():
    oc_values._ensure_headers(
        [f" {header}. " for header in EDI_HEADERS], EDI_HEADERS, "Sheet1"
    )


def test_a_wrong_header_names_the_column_and_both_values():
    actual = list(EDI_HEADERS)
    actual[1] = "Cot Sai"

    with pytest.raises(OCWorkbookError) as error:
        oc_values._ensure_headers(actual, EDI_HEADERS, "Sheet1")

    assert error.value.code == "OC_FILE_HEADERS_INVALID"
    assert "Cột 2" in error.value.errors[0]
    assert "Cot Sai" in error.value.errors[0]


def test_a_missing_trailing_column_is_reported_as_empty():
    with pytest.raises(OCWorkbookError) as error:
        oc_values._ensure_headers(list(EDI_HEADERS[:-1]), EDI_HEADERS, "Sheet1")

    assert "[trống]" in error.value.errors[-1]


def test_the_error_list_stops_at_eight_so_the_dialog_stays_readable():
    with pytest.raises(OCWorkbookError) as error:
        oc_values._ensure_headers(
            ["sai"] * len(EDI_HEADERS), EDI_HEADERS, "Sheet1"
        )

    assert len(error.value.errors) == 8


# --- sheet tham chiếu của form cũ --------------------------------------


def _legacy_workbook(rows):
    workbook = Workbook()
    sheet = workbook.active
    sheet.title = "THONG TIN"
    sheet.append(["Factory", "Buyer", "", "", "Destination", "Market"])
    for row in rows:
        sheet.append(list(row))
    return workbook


def test_a_form_without_the_reference_sheet_is_reported():
    workbook = Workbook()
    workbook.active.title = "FORM"

    with pytest.raises(OCWorkbookError) as error:
        oc_values._lookup_lists(workbook)

    assert error.value.code == "OC_TEMPLATE_SHEET_MISSING"


def test_the_buyer_list_stops_at_the_po_type_marker():
    workbook = _legacy_workbook(
        [
            ["PSHK", "J.LINDEBERG", "", "", "VIETNAM", "ASIA"],
            ["", "TRUEWERK", "", "", "", ""],
            ["", "PO Type", "", "", "", ""],
            ["", "KHONG TINH", "", "", "", ""],
        ]
    )

    factories, buyers, countries = oc_values._lookup_lists(workbook)

    assert "pshk" in factories
    assert buyers == {"j.lindeberg", "truewerk"}
    assert countries["vietnam"] == ("VIETNAM", "ASIA")
    workbook.close()


def test_blank_reference_rows_are_skipped():
    workbook = _legacy_workbook(
        [
            ["", "", "", "", "", ""],
            ["PSHK", "J.LINDEBERG", "", "", "VIETNAM", "ASIA"],
        ]
    )

    factories, buyers, _countries = oc_values._lookup_lists(workbook)

    assert factories == {"pshk"}
    assert buyers == {"j.lindeberg"}
    workbook.close()


def test_the_workbook_stream_helper_reads_what_was_written(tmp_path):
    path = _xlsx(tmp_path)

    assert path.read_bytes().startswith(b"PK")
    with ZipFile(BytesIO(path.read_bytes())) as archive:
        assert "[Content_Types].xml" in archive.namelist()
