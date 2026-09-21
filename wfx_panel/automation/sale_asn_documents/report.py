"""Tải Excel từ Report Viewer và nhận diện đúng loại workbook.

Tải bằng fetch bất đồng bộ ngay trong Report Viewer với phiên đang đăng nhập,
không bấm menu Excel qua ``window.open`` — Chrome có thể chặn im lặng hoặc tạo
download trễ/trùng. Trong lúc SSRS tạo file, worker poll theo lát có checkpoint
hủy chứ không chặn một ``request.get`` tới hết timeout."""

from __future__ import annotations

import base64
from pathlib import Path
from typing import Any
from urllib.parse import urljoin

from openpyxl import load_workbook

from wfx_panel.automation._common import (
    Callable,
    Frame,
    Page,
    PlaywrightError,
    PlaywrightTimeoutError,
    _wait,
    _write_log,
    time,
)
from wfx_panel.automation.runtime import cancellation_deferred
from wfx_panel.automation.sale_asn_documents.constants import (
    _REPORT_FETCH_CHUNK_JS,
    _REPORT_FETCH_CLEANUP_JS,
    _REPORT_FETCH_START_JS,
    _REPORT_FETCH_STATE_JS,
    BUYER_INVOICE_SELECTOR,
    DOCUMENTS_FRAME_TIMEOUT_SECONDS,
    PACKING_LIST_SELECTOR,
    REPORT_DOWNLOAD_CHUNK_BYTES,
    REPORT_DOWNLOAD_MAX_ATTEMPTS,
    REPORT_DOWNLOAD_POLL_MS,
    REPORT_DOWNLOAD_PROGRESS_INTERVAL_SECONDS,
    REPORT_DOWNLOAD_RETRY_DELAY_MS,
    REPORT_DOWNLOAD_START_TIMEOUT_SECONDS,
    REPORT_EXPORT_SELECTOR,
    REPORT_READY_TIMEOUT_SECONDS,
)


def _find_frame_with(
    context: Any,
    selectors: tuple[str, ...],
    *,
    timeout_s: float,
    visible: bool = False,
) -> tuple[Page, Frame]:
    deadline = time.monotonic() + timeout_s
    while time.monotonic() < deadline:
        for page in reversed(context.pages):
            for frame in reversed(page.frames):
                try:
                    matched = True
                    for selector in selectors:
                        locator = frame.locator(selector)
                        if not locator.count() or (
                            visible and not locator.first.is_visible()
                        ):
                            matched = False
                            break
                    if matched:
                        return page, frame
                except PlaywrightError:
                    continue
        first_page = context.pages[0] if context.pages else None
        if first_page is not None:
            _wait(first_page, 150)
    raise PlaywrightTimeoutError(
        "Không tìm thấy màn hình chứa: " + ", ".join(selectors)
    )


def _mark_report_frames(context: Any) -> list[tuple[Frame, str]]:
    snapshots: list[tuple[Frame, str]] = []
    for page in context.pages:
        for frame in page.frames:
            try:
                if not frame.locator(REPORT_EXPORT_SELECTOR).count():
                    continue
                marker = f"asn-report-{time.monotonic_ns()}"
                frame.evaluate(
                    "marker => { window.__wfxAsnReportMarker = marker; }",
                    marker,
                )
                snapshots.append((frame, marker))
            except PlaywrightError:
                continue
    return snapshots


def _report_frame_is_new(
    frame: Frame,
    snapshots: list[tuple[Frame, str]],
) -> bool:
    old = next((item for item in snapshots if item[0] == frame), None)
    if old is None:
        return True
    try:
        marker = frame.evaluate("() => window.__wfxAsnReportMarker || ''")
        return marker != old[1]
    except PlaywrightError:
        return True


def _report_export_url(frame: Frame) -> str:
    """Trả URL export chỉ khi Report Viewer đã khởi tạo hoàn chỉnh.

    Toolbar có thể sáng trước khi SSRS gán ``ExportUrlBase``. Nếu click Excel
    trong khoảng này, ``ExportReport`` chỉ trả ``false`` và không tạo request
    download dù nội dung report đã hiện đầy đủ.
    """
    try:
        return str(
            frame.evaluate(
                """() => {
                    const viewer = window.$find?.('rptCustomReportViewer');
                    const internal = viewer?._getInternalViewer?.();
                    return internal?.ExportUrlBase || '';
                }"""
            )
            or ""
        ).strip()
    except PlaywrightError:
        return ""


def _wait_report_ready(
    context: Any,
    snapshots: list[tuple[Frame, str]],
    timeout_s: float = REPORT_READY_TIMEOUT_SECONDS,
) -> tuple[Page, Frame]:
    deadline = time.monotonic() + timeout_s
    stable_since = 0.0
    candidate_key: tuple[int, int] | None = None
    while time.monotonic() < deadline:
        try:
            page, frame = _find_frame_with(
                context,
                (REPORT_EXPORT_SELECTOR,),
                timeout_s=0.2,
                visible=True,
            )
            if not _report_frame_is_new(frame, snapshots):
                if context.pages:
                    _wait(context.pages[0], 150)
                continue
            export = frame.locator(REPORT_EXPORT_SELECTOR).first
            loading = frame.locator(
                '[id*="AsyncWait"][style*="display: block"], '
                '[id*="AsyncWait"][aria-hidden="false"], [aria-busy="true"]'
            )
            has_loading = any(
                loading.nth(index).is_visible()
                for index in range(loading.count())
            )
            key = (id(page), id(frame))
            now = time.monotonic()
            export_url = _report_export_url(frame)
            if export.is_enabled() and not has_loading and export_url:
                if key != candidate_key:
                    candidate_key = key
                    stable_since = now
                elif now - stable_since >= 0.8:
                    return page, frame
            else:
                candidate_key = None
                stable_since = 0.0
        except (PlaywrightError, PlaywrightTimeoutError):
            candidate_key = None
            stable_since = 0.0
        if context.pages:
            _wait(context.pages[0], 150)
    raise PlaywrightTimeoutError("Report Viewer chưa load xong.")


def _download_report_excel(
    _context: Any,
    report_frame: Frame,
    target: Path,
    label: str,
    log: Callable[[str], None],
) -> None:
    _write_log(log, f"[SALE ASN DOCS] Đang tải Excel: {label}...")
    deadline = time.monotonic() + REPORT_DOWNLOAD_START_TIMEOUT_SECONDS
    response_body = b""
    last_status: int | str = "unknown"
    attempt = 0
    while (
        attempt < REPORT_DOWNLOAD_MAX_ATTEMPTS
        and time.monotonic() < deadline
    ):
        attempt += 1
        export_base = _report_export_url(report_frame)
        if not export_base:
            raise PlaywrightTimeoutError(f"WFX chưa tạo link download {label}.")
        export_url = urljoin(
            report_frame.url,
            f"{export_base}EXCELOPENXML",
        )
        report_frame.evaluate(_REPORT_FETCH_START_JS, export_url)
        fetch_started = time.monotonic()
        next_progress = (
            fetch_started + REPORT_DOWNLOAD_PROGRESS_INTERVAL_SECONDS
        )
        try:
            while time.monotonic() < deadline:
                state = report_frame.evaluate(_REPORT_FETCH_STATE_JS)
                if state.get("done"):
                    last_status = state.get("status") or state.get("error") or "unknown"
                    if state.get("ok") and state.get("prefix") == "PK":
                        chunks: list[bytes] = []
                        size = max(0, int(state.get("size") or 0))
                        for offset in range(0, size, REPORT_DOWNLOAD_CHUNK_BYTES):
                            # _wait(..., 0) chỉ chạy checkpoint, nhờ đó Stop vẫn
                            # phản hồi khi workbook lớn đang được chuyển từ frame.
                            _wait(report_frame, 0)
                            encoded = report_frame.evaluate(
                                _REPORT_FETCH_CHUNK_JS,
                                {
                                    "offset": offset,
                                    "size": REPORT_DOWNLOAD_CHUNK_BYTES,
                                },
                            )
                            chunks.append(base64.b64decode(encoded))
                        response_body = b"".join(chunks)
                    break
                now = time.monotonic()
                if now >= next_progress:
                    waited = int(now - fetch_started)
                    _write_log(
                        log,
                        f"[SALE ASN DOCS] WFX vẫn đang tạo {label}; đã chờ {waited} giây...",
                    )
                    next_progress = (
                        now + REPORT_DOWNLOAD_PROGRESS_INTERVAL_SECONDS
                    )
                remaining_ms = max(1, int((deadline - now) * 1_000))
                _wait(report_frame, min(REPORT_DOWNLOAD_POLL_MS, remaining_ms))
        finally:
            try:
                report_frame.evaluate(_REPORT_FETCH_CLEANUP_JS)
            except PlaywrightError:
                pass
        if response_body.startswith(b"PK"):
            break
        if time.monotonic() >= deadline:
            break
        _write_log(
            log,
            f"[SALE ASN DOCS] {label} chưa trả file Excel ở lượt {attempt}; "
            "đang chờ WFX hoàn tất rồi tải lại cùng report...",
        )
        _wait(report_frame, REPORT_DOWNLOAD_RETRY_DELAY_MS)
    else:
        response_body = b""
    if not response_body.startswith(b"PK"):
        raise RuntimeError(
            f"WFX trả file {label} không hợp lệ (HTTP {last_status})."
        )
    with cancellation_deferred():
        target.parent.mkdir(parents=True, exist_ok=True)
        target.write_bytes(response_body)
    if not target.is_file() or target.stat().st_size <= 0:
        raise RuntimeError(f"File {label} tải về bị rỗng.")
    _validate_report_kind(target, label)
    _write_log(log, f"[SALE ASN DOCS] Đã tải đúng {label}.")


def _report_workbook_kind(path: Path) -> str:
    """Nhận diện rõ report vừa tải; không suy đoán nếu WFX đổi template."""
    try:
        workbook = load_workbook(path, read_only=True, data_only=True)
    except Exception as error:
        raise RuntimeError(f"File report Excel không đọc được: {error}") from error
    try:
        text_parts: list[str] = []
        for sheet in workbook.worksheets:
            text_parts.append(sheet.title)
            for row_index, row in enumerate(sheet.iter_rows(values_only=True)):
                if row_index >= 20:
                    break
                text_parts.extend(
                    str(value)
                    for value in row[:20]
                    if value not in (None, "")
                )
        normalized = " ".join(" ".join(text_parts).casefold().split())
        if "packing list" in normalized:
            return "packing"
        if "commercial invoice" in normalized or "buyer invoice" in normalized:
            return "invoice"
        return ""
    finally:
        workbook.close()


def _validate_report_kind(path: Path, label: str) -> None:
    expected = "packing" if label == "Packing List" else "invoice"
    actual = _report_workbook_kind(path)
    if actual and actual != expected:
        actual_label = "Packing List" if actual == "packing" else "Buyer Invoice"
        raise RuntimeError(
            f"WFX trả nhầm {actual_label} khi đang tải {label}."
        )


def _documents_frame(
    context: Any,
    timeout_s: float = DOCUMENTS_FRAME_TIMEOUT_SECONDS,
) -> tuple[Page, Frame]:
    return _find_frame_with(
        context,
        (PACKING_LIST_SELECTOR, BUYER_INVOICE_SELECTOR),
        timeout_s=timeout_s,
        visible=True,
    )


def _restore_documents_screen(
    context: Any,
    report_page: Page,
    report_frame: Frame,
    docs_url: str,
) -> tuple[Page, Frame]:
    try:
        return _documents_frame(context, timeout_s=1)
    except PlaywrightTimeoutError:
        pass
    try:
        report_frame.evaluate("history.back()")
    except PlaywrightError:
        try:
            report_frame.goto(docs_url, wait_until="domcontentloaded", timeout=15_000)
        except PlaywrightError:
            try:
                report_page.go_back(wait_until="domcontentloaded", timeout=15_000)
            except PlaywrightError:
                pass
    return _documents_frame(context, timeout_s=DOCUMENTS_FRAME_TIMEOUT_SECONDS)


def _close_sale_asn_document_popups(
    context: Any,
    existing_page_ids: set[int],
    log: Callable[[str], None],
) -> None:
    """Đóng các popup sinh từ Docs sau khi đã lưu xong file ghép.

    Chỉ đóng Page không có trước lúc click Docs để không đụng cửa sổ WFX mà
    người dùng đã mở sẵn. Khi WFX tái sử dụng tab List hiện tại, tab đó cũng
    được giữ nguyên.
    """
    closed = 0
    for page in reversed(context.pages):
        if id(page) in existing_page_ids:
            continue
        try:
            if not page.is_closed():
                page.close(run_before_unload=False)
                closed += 1
        except PlaywrightError:
            continue
    if closed:
        _write_log(log, f"[SALE ASN DOCS] Đã đóng {closed} cửa sổ Docs/report.")
