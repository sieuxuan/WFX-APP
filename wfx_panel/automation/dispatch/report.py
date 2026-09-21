"""Mở report BuyerDispatchOrder_Invoice, xuất Excel rồi lưu lại thành XLSX."""

from __future__ import annotations

import re
from pathlib import Path

from openpyxl import load_workbook

from wfx_panel.automation._common import (
    Any,
    Callable,
    Page,
    PlaywrightError,
    PlaywrightTimeoutError,
    _click,
    _first_line,
    _wait,
    _write_log,
    time,
)
from wfx_panel.automation.dispatch.constants import (
    REPORT_DOC_NO_SELECTOR,
    REPORT_EXCEL_SELECTOR,
    REPORT_EXPORT_IMAGE_SELECTOR,
    REPORT_EXPORT_LINK_SELECTOR,
    REPORT_EXPORT_MENU_SELECTOR,
    REPORT_TIMEOUT_SECONDS,
    REPORT_URL,
    REPORT_VIEW_SELECTOR,
    _emit_progress,
)
from wfx_panel.automation.dispatch.status import DispatchFlowError
from wfx_panel.automation.runtime import (
    cancellation_deferred,
    checkpoint,
    save_native_download,
    snapshot_downloads,
)


def reload_dispatch_workbook(source: Path, target: Path) -> None:
    """Mở và ghi lại report thành XLSX sạch để WFX import ổn định."""
    workbook = None
    try:
        workbook = load_workbook(source, data_only=False, keep_links=False)
        if not workbook.sheetnames:
            raise ValueError("Workbook không có sheet.")
        if not any(
            sheet.max_row > 0 and sheet.max_column > 0
            for sheet in workbook.worksheets
        ):
            raise ValueError("Workbook không có dữ liệu.")
        calculation = getattr(workbook, "calculation", None)
        if calculation is not None:
            calculation.fullCalcOnLoad = True
            calculation.forceFullCalc = True
            calculation.calcMode = "auto"
        target.parent.mkdir(parents=True, exist_ok=True)
        workbook.save(target)
    except Exception as error:
        raise DispatchFlowError(
            "GDN_WORKBOOK_RELOAD_FAILED",
            "Không thể reload file report thành XLSX để import.",
            errors=[f"{type(error).__name__}: {_first_line(error)}"],
        ) from error
    finally:
        if workbook is not None:
            workbook.close()
    if not target.is_file() or target.stat().st_size <= 0:
        raise DispatchFlowError(
            "GDN_WORKBOOK_RELOAD_FAILED",
            "File XLSX sau khi reload bị rỗng.",
        )


def _wait_report_ready(report_page: Page) -> None:
    deadline = time.monotonic() + REPORT_TIMEOUT_SECONDS
    last_report_text = ""
    while time.monotonic() < deadline:
        checkpoint()
        try:
            image = report_page.locator(REPORT_EXPORT_IMAGE_SELECTOR)
            if image.count():
                source = str(image.first.get_attribute("src") or "")
                async_wait = report_page.locator(
                    "#rptCustomReportViewer_AsyncWait"
                )
                loading = async_wait.count() and async_wait.first.is_visible()
                if source and "disabled" not in source.casefold() and not loading:
                    return
            report = report_page.locator("#rptCustomReportViewer_ctl09")
            if report.count():
                last_report_text = " ".join(
                    (report.first.inner_text(timeout=1_000) or "").split()
                )
                if re.search(r"no\s+(data|rows)|không\s+có\s+dữ\s+liệu", last_report_text, re.I):
                    raise DispatchFlowError(
                        "GDN_REPORT_EMPTY",
                        "Report không có dữ liệu cho Invoice GRN đã nhập.",
                    )
        except DispatchFlowError:
            raise
        except PlaywrightError:
            pass
        _wait(report_page, 150)
    suffix = f" Chi tiết: {last_report_text[:300]}" if last_report_text else ""
    raise DispatchFlowError(
        "GDN_REPORT_NOT_READY",
        f"Report Buyer Dispatch chưa load xong.{suffix}",
    )


def _download_report(
    report_page: Page,
    target: Path,
    log: Callable[[str], None],
) -> None:
    downloads: list[Any] = []

    def receive(download: Any) -> None:
        downloads.append(download)

    report_page.on("download", receive)
    try:
        export = report_page.locator(REPORT_EXPORT_LINK_SELECTOR)
        if not export.count() or not export.first.is_visible():
            raise DispatchFlowError(
                "GDN_REPORT_NOT_READY",
                "Nút Export của report chưa sẵn sàng.",
            )
        _click(export.first)
        menu = report_page.locator(REPORT_EXPORT_MENU_SELECTOR)
        menu.wait_for(state="visible", timeout=5_000)
        excel = report_page.locator(REPORT_EXCEL_SELECTOR)
        if excel.count() != 1 or not excel.first.is_visible():
            raise DispatchFlowError(
                "GDN_REPORT_NOT_READY",
                "Lựa chọn Excel trong menu Export chưa sẵn sàng.",
            )
        _write_log(log, "[GDN] Đang export report sang Excel...")
        downloads_before_click = snapshot_downloads()
        _click(excel.first)
        deadline = time.monotonic() + 90
        while time.monotonic() < deadline and not downloads:
            _wait(report_page, 100)
        if not downloads:
            raise DispatchFlowError(
                "GDN_REPORT_DOWNLOAD_FAILED",
                "WFX không bắt đầu tải file Excel của report.",
            )
        with cancellation_deferred():
            target.parent.mkdir(parents=True, exist_ok=True)
            save_native_download(
                downloads[0],
                target,
                downloads_before_click,
            )
        if not target.is_file() or target.stat().st_size <= 0:
            raise DispatchFlowError(
                "GDN_REPORT_DOWNLOAD_FAILED",
                "File Excel tải từ report bị rỗng.",
            )
    finally:
        try:
            report_page.remove_listener("download", receive)
        except Exception:
            pass


def _prepare_dispatch_workbook(
    context: Any,
    invoice: str,
    temporary: Path,
    log: Callable[[str], None],
    progress: Callable[..., None] | None = None,
) -> Path:
    report_page: Page | None = None
    try:
        report_page = context.new_page()
        report_page.goto(
            REPORT_URL,
            wait_until="domcontentloaded",
            timeout=30_000,
        )
        doc_no = report_page.locator(REPORT_DOC_NO_SELECTOR)
        doc_no.wait_for(state="visible", timeout=20_000)
        doc_no.fill(invoice)
        if doc_no.input_value() != invoice:
            raise DispatchFlowError(
                "GDN_REPORT_NOT_READY",
                "WFX chưa xác nhận Doc No. trên report.",
            )
        _write_log(log, "[GDN] Đã điền Invoice GRN vào Doc No.")
        _click(report_page.locator(REPORT_VIEW_SELECTOR))
        _write_log(log, "[GDN] Đang chờ report Buyer Dispatch load...")
        _emit_progress(
            progress,
            "report",
            "Đang tải báo cáo Buyer Dispatch…",
            1,
        )
        _wait_report_ready(report_page)
        _emit_progress(
            progress,
            "download",
            "Báo cáo đã sẵn sàng · đang tải Excel…",
            2,
        )
        raw_report = temporary / "BuyerDispatchOrder_Invoice.download.xlsx"
        _download_report(report_page, raw_report, log)
        _emit_progress(
            progress,
            "workbook",
            "Đang chuẩn hóa workbook XLSX…",
            3,
        )
        upload_path = temporary / "BuyerDispatchOrder_Invoice.reload.xlsx"
        reload_dispatch_workbook(raw_report, upload_path)
        _write_log(log, "[GDN] Đã reload và save report thành XLSX.")
        return upload_path
    except DispatchFlowError:
        raise
    except PlaywrightTimeoutError as error:
        raise DispatchFlowError(
            "GDN_REPORT_NOT_READY",
            "Report Buyer Dispatch chưa sẵn sàng.",
            errors=[_first_line(error)],
        ) from error
    except PlaywrightError as error:
        raise DispatchFlowError(
            "GDN_REPORT_DOWNLOAD_FAILED",
            "Không thể tải report Buyer Dispatch từ WFX.",
            errors=[_first_line(error)],
        ) from error
    finally:
        if report_page is not None:
            try:
                report_page.close()
            except PlaywrightError:
                pass
