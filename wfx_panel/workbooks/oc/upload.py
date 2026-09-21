"""Dựng workbook tạm chỉ có `Sheet1` rồi xác minh lại trước khi giao cho EDI."""

from __future__ import annotations

from decimal import Decimal
from pathlib import Path
from typing import Any
from zipfile import BadZipFile

from openpyxl import Workbook, load_workbook
from openpyxl.styles import Alignment, Font, PatternFill

from wfx_panel.workbooks.oc.rows import _new_rows, _revise_rows
from wfx_panel.workbooks.oc.schema import (
    DATE_HEADERS,
    EDI_HEADERS,
    OCWorkbookError,
    PreparedOCUpload,
)
from wfx_panel.workbooks.oc.values import (
    _ensure_headers,
    _ensure_input_values_only,
    _normalise_text,
    _validate_xlsx_archive,
)


def _excel_value(value: Any) -> Any:
    if isinstance(value, Decimal):
        return int(value) if value == value.to_integral_value() else float(value)
    return value


def _write_static_workbook(rows: list[list[Any]], output_path: Path) -> None:
    output_path.parent.mkdir(parents=True, exist_ok=True)
    workbook = Workbook()
    sheet = workbook.active
    sheet.title = "Sheet1"
    sheet.append(list(EDI_HEADERS))
    header_fill = PatternFill("solid", fgColor="FFFF00")
    for cell in sheet[1]:
        cell.fill = header_fill
        cell.font = Font(bold=True)
        cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
    for source in rows:
        sheet.append([_excel_value(value) for value in source])
    date_indexes = [EDI_HEADERS.index(header) + 1 for header in DATE_HEADERS]
    for column in date_indexes:
        for row_number in range(2, sheet.max_row + 1):
            sheet.cell(row_number, column).number_format = "dd-mm-yyyy"
    sheet.freeze_panes = "A2"
    sheet.auto_filter.ref = f"A1:AY{sheet.max_row}"
    sheet.row_dimensions[1].height = 42
    workbook.save(output_path)


def _verify_static_output(path: Path, expected_rows: int) -> None:
    workbook = load_workbook(path, data_only=False, read_only=True)
    try:
        if workbook.sheetnames != ["Sheet1"]:
            raise OCWorkbookError(
                "OC_OUTPUT_INVALID",
                "File EDI sinh ra phải chỉ có một sheet Sheet1.",
            )
        sheet = workbook["Sheet1"]
        _ensure_headers(
            [sheet.cell(1, column).value for column in range(1, len(EDI_HEADERS) + 1)],
            EDI_HEADERS,
            "Sheet1",
        )
        if sheet.max_row - 1 != expected_rows:
            raise OCWorkbookError(
                "OC_OUTPUT_INVALID",
                "Số dòng trong file EDI sinh ra không khớp dữ liệu nguồn.",
            )
        for row in sheet.iter_rows(min_row=2, max_col=len(EDI_HEADERS)):
            for cell in row:
                if cell.data_type == "f":
                    raise OCWorkbookError(
                        "OC_OUTPUT_INVALID",
                        "File EDI sinh ra vẫn còn công thức Excel.",
                    )
    finally:
        workbook.close()


def prepare_oc_workbook(
    input_path: str | Path,
    mode: str,
    output_path: str | Path,
) -> PreparedOCUpload:
    """Validate source then create a value-only 51-column EDI workbook."""
    source = Path(input_path).expanduser().resolve()
    target = Path(output_path).expanduser().resolve()
    selected_mode = str(mode or "").strip().casefold()
    if selected_mode not in {"new", "revise"}:
        raise OCWorkbookError(
            "OC_MODE_INVALID",
            "Chế độ Upload OC phải là New hoặc Revise.",
        )
    _validate_xlsx_archive(source)
    _ensure_input_values_only(source, selected_mode)
    try:
        workbook = load_workbook(source, data_only=True, read_only=False)
    except (OSError, ValueError, BadZipFile) as error:
        raise OCWorkbookError(
            "OC_FILE_INVALID",
            f"Không đọc được workbook: {type(error).__name__}: {error}",
        ) from error
    try:
        if selected_mode == "new":
            buyer, rows, warnings = _new_rows(workbook)
        else:
            buyer, rows, warnings = _revise_rows(workbook)
    finally:
        workbook.close()
    _write_static_workbook(rows, target)
    _verify_static_output(target, len(rows))
    indexes = {header: index for index, header in enumerate(EDI_HEADERS)}
    seasons = tuple(
        sorted(
            {
                _normalise_text(row[indexes["Season"]])
                for row in rows
                if _normalise_text(row[indexes["Season"]])
            },
            key=str.casefold,
        )
    )
    po_refs = {
        _normalise_text(row[indexes["Summary Buyer Order Ref"]]).casefold()
        for row in rows
        if _normalise_text(row[indexes["Summary Buyer Order Ref"]])
    }
    articles = {
        _normalise_text(row[indexes["Article"]]).casefold()
        for row in rows
        if _normalise_text(row[indexes["Article"]])
    }
    total_units = sum(
        (Decimal(str(row[indexes["Units"]])) for row in rows),
        Decimal(0),
    )
    return PreparedOCUpload(
        mode=selected_mode,
        buyer=buyer,
        row_count=len(rows),
        upload_path=target,
        seasons=seasons,
        po_count=len(po_refs),
        style_count=len(articles),
        total_units=_excel_value(total_units),
        warnings=warnings,
    )
