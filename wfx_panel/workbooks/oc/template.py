"""Sinh form `OC INPUT` cho người dùng nhập.

Sheet `REFERENCES` phải `veryHidden` và chỉ chứa nguồn dropdown; user không
được nhìn thấy hay sửa nó."""

from __future__ import annotations

from pathlib import Path
from typing import Any

from openpyxl import Workbook
from openpyxl.comments import Comment
from openpyxl.styles import Alignment, Font, PatternFill

from wfx_panel.workbooks.oc.schema import (
    BUYER_OPTIONS,
    DESTINATION_COUNTRY_MARKET,
    FACTORY_OPTIONS,
    INPUT_COMMENTS,
    INPUT_HEADERS,
    INPUT_SHEET_NAME,
    MAX_OC_ROWS,
    ORDER_TYPE_OPTIONS,
    PAYMENT_TERM_OPTIONS,
    PO_TYPE_OPTIONS,
    REFERENCE_SHEET_NAME,
    OCWorkbookError,
)
from wfx_panel.workbooks.oc.values import (
    _add_date_sequence_validation,
    _add_list_validation,
)


def write_oc_input_template(path: str | Path) -> Path:
    """Create the simplified one-header workbook shown to end users."""
    target = Path(path).expanduser().resolve()
    target.parent.mkdir(parents=True, exist_ok=True)
    workbook = Workbook()
    sheet = workbook.active
    sheet.title = INPUT_SHEET_NAME
    sheet.append(list(INPUT_HEADERS))
    required_fill = PatternFill("solid", fgColor="FFFF00")
    optional_fill = PatternFill("solid", fgColor="F4B183")
    optional_headers = {"PO Type (Zone)", "Extra Production %", "Buyer Lot No."}
    for cell in sheet[1]:
        cell.fill = optional_fill if cell.value in optional_headers else required_fill
        cell.font = Font(bold=True)
        cell.alignment = Alignment(
            horizontal="center",
            vertical="center",
            wrap_text=True,
        )
        comment = INPUT_COMMENTS.get(str(cell.value))
        if comment:
            cell.comment = Comment(comment, "WFX Smart")
    sheet.row_dimensions[1].height = 48
    sheet.freeze_panes = "A2"
    sheet.auto_filter.ref = "A1:X10001"
    widths = {
        "Buyer": 24,
        "Season": 12,
        "Order Type": 14,
        "Currency": 10,
        "Factory": 34,
        "Ship Under PO Ref": 28,
        "Article Code": 16,
        "Buyer Style Ref": 18,
        "Buyer PO Num": 28,
        "Summary Buyer Order Ref": 28,
        "Buyer Order Date": 17,
        "Buyer Delivery Date": 18,
        "Raw Material ETA Date": 20,
        "Payment Terms": 42,
        "Country of Final Destination": 25,
        "Color Code": 14,
        "Color Name": 20,
        "Size Code": 12,
        "Selling Price": 14,
        "Units": 11,
        "Internal Lot No.": 16,
        "PO Type (Zone)": 16,
        "Extra Production %": 18,
        "Buyer Lot No.": 18,
    }
    for index, header in enumerate(INPUT_HEADERS, start=1):
        sheet.column_dimensions[sheet.cell(1, index).column_letter].width = widths[header]
    for row_number in range(2, 202):
        for column in (11, 12, 13):
            sheet.cell(row_number, column).number_format = "dd-mm-yyyy"
        sheet.cell(row_number, 19).number_format = "0.00"
        sheet.cell(row_number, 20).number_format = "0"
        sheet.cell(row_number, 23).number_format = "0.00"

    references = workbook.create_sheet(REFERENCE_SHEET_NAME)
    reference_lists = (
        ("Buyer", BUYER_OPTIONS),
        ("Factory", FACTORY_OPTIONS),
        ("Order Type", ORDER_TYPE_OPTIONS),
        ("Currency", ("USD", "EUR", "GBP")),
        ("Country", tuple(DESTINATION_COUNTRY_MARKET)),
        ("PO Type (Zone)", PO_TYPE_OPTIONS),
        ("Payment Terms", PAYMENT_TERM_OPTIONS),
    )
    for column, (heading, options) in enumerate(reference_lists, start=1):
        references.cell(1, column, heading)
        for row, option in enumerate(options, start=2):
            references.cell(row, column, option)
        _add_list_validation(sheet, references, heading, column, len(options))
    _add_date_sequence_validation(sheet)
    references.sheet_state = "veryHidden"
    workbook.save(target)
    return target


def _nonempty_rows(sheet: Any, start_row: int, width: int) -> list[tuple[int, list[Any]]]:
    rows: list[tuple[int, list[Any]]] = []
    for row_number in range(start_row, min(sheet.max_row, MAX_OC_ROWS + start_row) + 1):
        values = [sheet.cell(row_number, column).value for column in range(1, width + 1)]
        if any(value not in (None, "") for value in values):
            rows.append((row_number, values))
    if sheet.max_row > MAX_OC_ROWS + start_row:
        for row_number in range(MAX_OC_ROWS + start_row + 1, sheet.max_row + 1):
            if any(
                sheet.cell(row_number, column).value not in (None, "")
                for column in range(1, width + 1)
            ):
                raise OCWorkbookError(
                    "OC_FILE_TOO_MANY_ROWS",
                    f"Upload OC hỗ trợ tối đa {MAX_OC_ROWS:,} dòng dữ liệu.",
                )
    return rows
