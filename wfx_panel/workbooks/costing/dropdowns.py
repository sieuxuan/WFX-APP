"""Data validation của form Costing: dropdown cho từng ô nhập.

Article Code/Article Name dùng công thức lookup an toàn để tên đổi theo mã;
khi chưa có cache Article Library thì vẫn phải cho nhập tay chứ không được
khoá ô lại.
"""

from __future__ import annotations

from collections.abc import Iterable, Mapping
from typing import Any

from openpyxl.utils import get_column_letter
from openpyxl.worksheet.datavalidation import DataValidation

from wfx_panel.workbooks.costing.schema import (
    FORM_COLUMNS,
    _article_lookup_options,
    _base_field_key,
    _CostingFormLayout,
    _excel_safe,
    _section_item_type,
    _text,
)


def _unique_values(values: Iterable[Any]) -> list[str]:
    output: list[str] = []
    seen: set[str] = set()
    for value in values:
        text = _text(value)
        folded = text.casefold()
        if not text or folded in seen:
            continue
        seen.add(folded)
        output.append(text)
    return output

def _split_wfx_multiselect(value: Any) -> list[str]:
    """Tách chuỗi WFX theo dấu phẩy nhưng giữ nguyên dấu phẩy trong ngoặc."""
    text = _text(value)
    if not text:
        return []
    output: list[str] = []
    start = 0
    depth = 0
    for index, character in enumerate(text):
        if character == "(":
            depth += 1
        elif character == ")" and depth:
            depth -= 1
        elif character in {",", "|"} and depth == 0:
            token = text[start:index].strip()
            if token:
                output.append(token)
            start = index + 1
    token = text[start:].strip()
    if token:
        output.append(token)
    return output

def _form_dropdown_options(
    document: Mapping[str, Any],
) -> dict[str, list[str]]:
    """Tạo lookup chung; option Material được gắn riêng theo từng item."""
    fields = list(document.get("fields") or ())

    def matching(field_key: str) -> list[Mapping[str, Any]]:
        wanted = field_key.casefold()
        return [
            field
            for field in fields
            if _base_field_key(field.get("field_key")) == wanted
        ]

    purchase_officers = matching("colPurchaseOfficer")
    return {
        "Color Dep.": _unique_values(
            [
                "[None]",
                "[Table]",
                "[Body Type]",
                "[Base Colors]",
            ]
        ),
        "Size Dep.": _unique_values(
            [
                "[None]",
                "[Table]",
                "[Body Type]",
            ]
        ),
        "Purchase Officer": _unique_values(
            [
                *(field.get("value") for field in purchase_officers),
                *(
                    option
                    for field in purchase_officers
                    for option in field.get("options") or ()
                ),
            ]
        ),
    }

def _add_form_dropdowns(
    ws: Any,
    document: Mapping[str, Any],
    *,
    last_row: int,
) -> int:
    """Gắn dropdown bằng lookup columns ẩn, không tạo sheet thứ ba."""
    option_sets = _form_dropdown_options(document)
    lookup_column = len(FORM_COLUMNS) + 1
    for label, options in option_sets.items():
        if not options:
            continue
        column_letter = get_column_letter(lookup_column)
        for row, option in enumerate(options, 2):
            ws.cell(row=row, column=lookup_column, value=_excel_safe(option))
        ws.column_dimensions[column_letter].hidden = True
        validation = DataValidation(
            type="list",
            formula1=f"=${column_letter}$2:${column_letter}${len(options) + 1}",
            allow_blank=True,
        )
        # Cho phép gõ mapping mới hoặc nhiều giá trị bằng dấu | khi lookup live
        # chưa có đủ màu/size. Dropdown vẫn là đường nhập mặc định.
        validation.showErrorMessage = False
        validation.promptTitle = "WFX Smart"
        validation.prompt = (
            "Chọn từ danh sách; có thể gõ nhiều giá trị, ngăn bằng dấu |."
        )
        validation.showInputMessage = True
        ws.add_data_validation(validation)
        target_column = FORM_COLUMNS.index(label) + 1
        target_letter = get_column_letter(target_column)
        validation.add(f"{target_letter}2:{target_letter}{last_row}")
        lookup_column += 1
    return lookup_column

def _add_item_option_dropdowns(
    ws: Any,
    document: Mapping[str, Any],
    row_by_item: Mapping[tuple[str, str], int],
    *,
    lookup_column: int,
) -> None:
    """Gắn dropdown Material Color/Size đúng option của từng Article row."""
    wanted_fields = {
        "colmaterialcolorlist": "Material Color",
        "colmaterialsizelist": "Material Size",
    }
    for field in document.get("fields") or ():
        base_key = _base_field_key(field.get("field_key"))
        label = wanted_fields.get(base_key)
        if label is None:
            continue
        row = row_by_item.get(
            (
                _text(field.get("section_key")).casefold(),
                _text(field.get("item_key")).casefold(),
            )
        )
        options = _unique_values(
            [
                *(field.get("options") or ()),
                *_split_wfx_multiselect(field.get("value")),
            ]
        )
        if row is None or not options:
            continue
        lookup_letter = get_column_letter(lookup_column)
        for option_row, option in enumerate(options, 2):
            ws.cell(
                row=option_row,
                column=lookup_column,
                value=_excel_safe(option),
            )
        ws.column_dimensions[lookup_letter].hidden = True
        validation = DataValidation(
            type="list",
            formula1=(
                f"=${lookup_letter}$2:"
                f"${lookup_letter}${len(options) + 1}"
            ),
            allow_blank=True,
        )
        validation.showErrorMessage = False
        validation.promptTitle = f"{label} của Article"
        validation.prompt = (
            "Chọn một giá trị; nhiều giá trị có thể nhập theo đúng chuỗi WFX."
        )
        validation.showInputMessage = True
        ws.add_data_validation(validation)
        target_letter = get_column_letter(FORM_COLUMNS.index(label) + 1)
        validation.add(f"{target_letter}{row}")
        lookup_column += 1

def _add_special_article_dropdowns(
    ws: Any,
    document: Mapping[str, Any],
    layout: _CostingFormLayout,
    *,
    lookup_column: int,
) -> int:
    """Gắn danh sách nhà máy/quy trình/chi phí cho từng block đặc biệt."""
    article_name_letter = get_column_letter(FORM_COLUMNS.index("Article Name") + 1)
    ordered_sections = sorted(
        document["sections"],
        key=lambda value: value["row_order"],
    )
    for section, rows in zip(ordered_sections, layout.section_rows, strict=True):
        if _section_item_type(section) != "cost_line":
            continue
        options = _unique_values(section.get("article_options") or ())
        if not options or not rows:
            continue
        lookup_letter = get_column_letter(lookup_column)
        for option_row, option in enumerate(options, 2):
            ws.cell(option_row, lookup_column, _excel_safe(option))
        ws.column_dimensions[lookup_letter].hidden = True
        validation = DataValidation(
            type="list",
            formula1=(
                f"=${lookup_letter}$2:"
                f"${lookup_letter}${len(options) + 1}"
            ),
            allow_blank=True,
        )
        validation.showErrorMessage = False
        validation.promptTitle = "Chọn từ WFX"
        validation.prompt = "Chọn tên đã quét; có thể gõ để tìm trong Excel."
        validation.showInputMessage = True
        ws.add_data_validation(validation)
        validation.add(
            f"{article_name_letter}{min(rows)}:{article_name_letter}{max(rows)}"
        )
        lookup_column += 1
    return lookup_column

def _add_material_article_dropdowns(
    ws: Any,
    document: Mapping[str, Any],
    layout: _CostingFormLayout,
    *,
    lookup_column: int,
) -> int:
    """Dropdown Article Code/Name từ thư viện server/cache."""
    ordered_sections = sorted(
        document["sections"],
        key=lambda value: value["row_order"],
    )
    validation_cache: dict[
        tuple[tuple[str, str], ...],
        tuple[DataValidation, DataValidation, str, str, int],
    ] = {}
    code_column = FORM_COLUMNS.index("Article Code") + 1
    name_column = FORM_COLUMNS.index("Article Name") + 1
    code_target_letter = get_column_letter(code_column)
    name_target_letter = get_column_letter(name_column)
    for section, rows in zip(ordered_sections, layout.section_rows, strict=True):
        if _section_item_type(section) == "cost_line" or not rows:
            continue
        pairs = _article_lookup_options(section.get("article_lookup_options"))
        if not pairs:
            codes = _unique_values(section.get("article_code_options") or ())
            names = _unique_values(section.get("article_name_options") or ())
            pairs = [
                {"article_code": code, "article_name": names[index]}
                for index, code in enumerate(codes)
                if index < len(names)
            ]
        if not pairs:
            continue
        signature = tuple(
            (option["article_code"], option["article_name"])
            for option in pairs
        )
        target_code_range = (
            f"{code_target_letter}{min(rows)}:"
            f"{code_target_letter}{max(rows)}"
        )
        target_name_range = (
            f"{name_target_letter}{min(rows)}:"
            f"{name_target_letter}{max(rows)}"
        )
        cached = validation_cache.get(signature)
        if cached is None:
            lookup_code_letter = get_column_letter(lookup_column)
            lookup_name_letter = get_column_letter(lookup_column + 1)
            for option_row, option in enumerate(pairs, 2):
                ws.cell(
                    option_row,
                    lookup_column,
                    _excel_safe(option["article_code"]),
                )
                ws.cell(
                    option_row,
                    lookup_column + 1,
                    _excel_safe(option["article_name"]),
                )
            ws.column_dimensions[lookup_code_letter].hidden = True
            ws.column_dimensions[lookup_name_letter].hidden = True
            last_option_row = len(pairs) + 1

            def validation_for(
                letter: str,
                label: str,
                option_last_row: int = last_option_row,
            ) -> DataValidation:
                validation = DataValidation(
                    type="list",
                    formula1=(
                        f"=${letter}$2:${letter}${option_last_row}"
                    ),
                    allow_blank=True,
                )
                validation.showErrorMessage = False
                validation.promptTitle = f"{label} từ thư viện"
                validation.prompt = (
                    "Danh sách tự đồng bộ từ server; vẫn có thể nhập tay."
                )
                validation.showInputMessage = True
                ws.add_data_validation(validation)
                return validation

            code_validation = validation_for(
                lookup_code_letter,
                "Article Code",
            )
            name_validation = validation_for(
                lookup_name_letter,
                "Article Name",
            )
            cached = (
                code_validation,
                name_validation,
                lookup_code_letter,
                lookup_name_letter,
                last_option_row,
            )
            validation_cache[signature] = cached
            lookup_column += 2
        (
            code_validation,
            name_validation,
            lookup_code_letter,
            lookup_name_letter,
            last_option_row,
        ) = cached
        code_validation.add(target_code_range)
        name_validation.add(target_name_range)
        available_codes = {
            option["article_code"].casefold() for option in pairs
        }
        for row in rows:
            current_code = _text(ws.cell(row, code_column).value)
            current_name = _text(ws.cell(row, name_column).value)
            if current_name and current_code.casefold() not in available_codes:
                continue
            ws.cell(row, name_column).value = (
                f'=IFERROR(INDEX(${lookup_name_letter}$2:'
                f'${lookup_name_letter}${last_option_row},'
                f'MATCH({code_target_letter}{row},'
                f'${lookup_code_letter}$2:'
                f'${lookup_code_letter}${last_option_row},0)),"")'
            )
    return lookup_column
