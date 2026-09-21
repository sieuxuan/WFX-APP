"""Sheet `Hướng dẫn` của workbook Costing.

Sheet này chỉ để đọc: nó giải thích từng cột và luật Action cho người dùng,
không tham gia round-trip dữ liệu.
"""

from __future__ import annotations

from collections.abc import Iterable, Mapping, Sequence
from typing import Any

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font
from openpyxl.utils import get_column_letter

from wfx_panel.workbooks.costing.schema import (
    _HEADER_FILL,
    _HEADER_FONT,
    _INPUT_FILL,
    _STRUCTURE_BORDER,
    _TITLE_FONT,
    CLEAR_MARKER,
    FORMAT_VERSION,
    GUIDE_SHEET,
    _excel_safe,
)


def _set_header(ws: Any, row: int, columns: Sequence[str]) -> None:
    for column, value in enumerate(columns, 1):
        cell = ws.cell(row=row, column=column, value=value)
        cell.fill = _HEADER_FILL
        cell.font = _HEADER_FONT
        cell.alignment = Alignment(vertical="center", wrap_text=True)
    ws.row_dimensions[row].height = 28
    ws.auto_filter.ref = (
        f"A{row}:{get_column_letter(len(columns))}{max(row, ws.max_row)}"
    )
    ws.freeze_panes = f"A{row + 1}"

def _finish_sheet(
    ws: Any,
    widths: Mapping[int, float],
    *,
    editable_columns: Iterable[int] = (),
) -> None:
    editable = set(editable_columns)
    for column, width in widths.items():
        ws.column_dimensions[get_column_letter(column)].width = width
    for row in ws.iter_rows(min_row=2):
        for cell in row:
            cell.alignment = Alignment(vertical="top", wrap_text=True)
            cell.border = _STRUCTURE_BORDER
            if cell.column in editable:
                cell.fill = _INPUT_FILL
    ws.sheet_view.showGridLines = False

def _write_guide(workbook: Workbook, document: Mapping[str, Any]) -> None:
    ws = workbook.create_sheet(GUIDE_SHEET)
    ws.merge_cells("A1:F1")
    ws["A1"] = "WFX Smart · Catalog Costing"
    ws["A1"].font = _TITLE_FONT
    ws["A1"].alignment = Alignment(vertical="center")
    ws.row_dimensions[1].height = 34
    rows = [
        ("Style Code", document["style_code"]),
        ("Style Name", document.get("style_name") or "—"),
        ("Trạng thái khi tải", document.get("cost_sheet_status") or "—"),
        ("Phiên bản file", f"v{FORMAT_VERSION}"),
        (
            "Cách dùng",
            "Mở sheet Costing và sửa các ô màu vàng. Dòng không dùng thì để trống.",
        ),
        (
            "Phạm vi",
            "Có thể tải file ở mọi trạng thái. Chỉ Cost Sheet đang Open mới "
            "được cập nhật lại lên WFX.",
        ),
        (
            "Các nhóm trong file",
            ", ".join(
                str(section.get("name") or section.get("section_key") or "")
                for section in document.get("sections") or ()
                if str(
                    section.get("name") or section.get("section_key") or ""
                ).strip()
            )
            or "—",
        ),
        (
            "Thêm nhiều dòng",
            "Điền lần lượt các dòng vàng trong đúng nhóm. Nếu cần thêm dòng, "
            "sao chép dòng vàng cuối và chèn trước nhóm tiếp theo.",
        ),
        (
            "Tách dòng màu/size",
            "Hai dòng liền nhau cùng Article Code sẽ được tách thành hai dòng "
            "riêng trên WFX.",
        ),
        (
            "Phối màu/size",
            "Mỗi dòng phối ghi Màu/Size vật tư => Màu/Size của style. Nhiều lựa "
            "chọn được ngăn bằng dấu |.",
        ),
        (
            "Cột công thức",
            "Cons. Qty. Incl. Waste = Cons. Qty. × (1 + Waste %/100); "
            "Value in (USD) = Rate × Cons. Qty. Incl. Waste. Hai cột đỏ chỉ đọc.",
        ),
        (
            "Purchase Officer",
            "Chọn từ danh sách nếu ô này trên WFX đang trống. App sẽ báo trước "
            "khi còn thiếu.",
        ),
        (
            "Cột Action",
            "Để trống = thêm mới hoặc cập nhật. Chọn DELETE chỉ khi muốn xóa dòng.",
        ),
        (
            "Xóa giá trị",
            f"Ô trống sẽ giữ nguyên dữ liệu cũ. Muốn xóa, ghi đúng {CLEAR_MARKER}.",
        ),
        (
            "An toàn",
            "App luôn cho xem trước thay đổi và chỉ lưu sau khi điền xong.",
        ),
        (
            "CM / Production / Indirect",
            "Chọn tên trong danh sách. CM và Indirect dùng USD. Production tự "
            "đặt Minutes = 1. Dòng không chọn tên sẽ không được thêm.",
        ),
    ]
    for row, (label, value) in enumerate(rows, 3):
        label_cell = ws.cell(row=row, column=1, value=label)
        label_cell.font = Font(bold=True)
        label_cell.alignment = Alignment(vertical="top", wrap_text=True)
        ws.cell(row=row, column=2, value=_excel_safe(value))
        ws.merge_cells(start_row=row, start_column=2, end_row=row, end_column=6)
        ws.cell(row=row, column=2).alignment = Alignment(
            horizontal="left",
            vertical="top",
            wrap_text=True,
        )
        ws.row_dimensions[row].height = 42 if row >= 6 else 24
        if label == "Cách dùng":
            ws.cell(row=row, column=2).fill = _INPUT_FILL
    for column, width in enumerate((23, 18, 18, 18, 18, 18), 1):
        ws.column_dimensions[get_column_letter(column)].width = width
    ws.sheet_view.showGridLines = False
