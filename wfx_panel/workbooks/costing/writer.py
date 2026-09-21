"""Sinh file Costing: sheet Hướng dẫn, form nhập, dropdown và công thức.

Workbook có đúng hai sheet — `Hướng dẫn` và `Costing`. Chỉ round-trip field
item `editable=true`; hai cột đỏ công thức là chỉ-đọc."""

from __future__ import annotations

from collections.abc import Mapping, Sequence
from pathlib import Path
from typing import Any

from openpyxl import Workbook
from openpyxl.comments import Comment
from openpyxl.styles import Alignment, Font
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.datavalidation import DataValidation

from wfx_panel.workbooks.costing.dropdowns import (  # noqa: F401
    _add_form_dropdowns,
    _add_item_option_dropdowns,
    _add_material_article_dropdowns,
    _add_special_article_dropdowns,
    _base_field_key,
    _form_dropdown_options,
    _split_wfx_multiselect,
    _unique_values,
)
from wfx_panel.workbooks.costing.guide import (  # noqa: F401
    _finish_sheet,
    _set_header,
    _write_guide,
)
from wfx_panel.workbooks.costing.schema import (
    _HEADER_FONT,
    _READ_ONLY_FILL,
    _READ_ONLY_HEADER_FILL,
    _SECTION_BORDER,
    _SUBHEADER_FILL,
    _TEMPLATE_INPUT_FILL,
    FORM_BASE_COLUMNS,
    FORM_COLUMNS,
    FORM_SHEET,
    FORM_TECH_COLUMNS,
    STANDARD_ITEM_FIELDS,
    CostingWorkbookError,
    _CostingFormLayout,
    _excel_safe,
    _form_field_index,
    _form_field_key,
    _form_items_by_section,
    _preflight_path,
    _section_item_type,
    _standard_section_token,
    _template_row_count,
    _text,
    workbook_document,
)


def _form_row_values(
    section: Mapping[str, Any],
    item: Mapping[str, Any] | None,
    row_order: int,
    field_by_item: Mapping[tuple[str, str, str], Mapping[str, Any]],
) -> list[Any]:
    section_key = str(section["section_key"])
    item_key = "" if item is None else str(item["item_key"])
    values: list[Any] = [
        section["name"],
        "" if item is None or item["action"] == "UPSERT" else item["action"],
        "" if item is None else item["article_code"],
        "" if item is None else item["article_name"],
    ]
    for definition in STANDARD_ITEM_FIELDS:
        field_key = _form_field_key(definition, section)
        field = field_by_item.get(
            (
                section_key.casefold(),
                item_key.casefold(),
                field_key.casefold(),
            )
        )
        value = "" if field is None else field["value"]
        token = _standard_section_token(section)
        if definition["label"] == "Minutes" and token == "productioncosts":
            value = 1
        elif definition["label"] == "Curr." and token in {
            "cmcosts",
            "indirectcosts",
        }:
            value = "USD"
        values.append(value)
    values.extend([section_key, item_key, row_order, _section_item_type(section)])
    return values


def _write_form_data_rows(
    ws: Any,
    document: Mapping[str, Any],
) -> _CostingFormLayout:
    field_by_item = _form_field_index(document)
    items_by_section = _form_items_by_section(document)
    section_rows: list[list[int]] = []
    template_rows: list[tuple[int, str]] = []
    row_by_item: dict[tuple[str, str], int] = {}
    output_row = 2
    for section in sorted(
        document["sections"],
        key=lambda value: value["row_order"],
    ):
        section_key = str(section["section_key"])
        section_items = items_by_section.get(section_key.casefold(), [])
        highest_row_order = max(
            (int(item["row_order"]) for item in section_items),
            default=0,
        )
        current_section_rows: list[int] = []
        rows_to_write = [
            *section_items,
            *([None] * _template_row_count(section)),
        ]
        for template_index, item in enumerate(rows_to_write, 1):
            is_template = item is None
            row_order = (
                highest_row_order + template_index - len(section_items)
                if is_template
                else int(item["row_order"])
            )
            for column, value in enumerate(
                _form_row_values(
                    section,
                    item,
                    row_order,
                    field_by_item,
                ),
                1,
            ):
                ws.cell(
                    row=output_row,
                    column=column,
                    value=_excel_safe(value),
                )
            if is_template:
                template_rows.append((output_row, str(section["name"])))
            else:
                row_by_item[
                    (section_key.casefold(), str(item["item_key"]).casefold())
                ] = output_row
            current_section_rows.append(output_row)
            output_row += 1
        section_rows.append(current_section_rows)
    return _CostingFormLayout(section_rows, template_rows, row_by_item)


def _form_column_widths() -> dict[int, int]:
    widths = {1: 24, 2: 12, 3: 20, 4: 30}
    special_widths = {
        "Material Color": 32,
        "Material Size": 28,
        "Color Mapping": 44,
        "Size Mapping": 44,
        "Cons. Qty. Incl. Waste": 24,
        "Value in (USD)": 20,
        "Remarks": 28,
        "Supplier": 25,
        "Placement": 20,
        "Purchase Officer": 20,
        "Shrinkage %(LxW)": 20,
    }
    for offset, definition in enumerate(
        STANDARD_ITEM_FIELDS,
        len(FORM_BASE_COLUMNS) + 1,
    ):
        widths[offset] = special_widths.get(str(definition["label"]), 16)
    return widths


def _style_formula_columns(ws: Any, last_row: int) -> set[int]:
    read_only_definitions = [
        definition
        for definition in STANDARD_ITEM_FIELDS
        if definition.get("read_only")
    ]
    read_only_columns: set[int] = set()
    for definition in read_only_definitions:
        column = FORM_COLUMNS.index(str(definition["label"])) + 1
        read_only_columns.add(column)
        header = ws.cell(row=1, column=column)
        header.fill = _READ_ONLY_HEADER_FILL
        header.font = _HEADER_FONT
        for row in range(2, last_row + 1):
            cell = ws.cell(row=row, column=column)
            cell.fill = _READ_ONLY_FILL
            cell.font = Font(color="991B1B")
        formula = (
            "Cons. Qty. × (1 + Waste %/100)"
            if definition["field_key"] == "colConsPlusWastageQty"
            else "Rate × Cons. Qty. Incl. Waste"
        )
        header.comment = Comment(
            f"Cột công thức WFX, chỉ đọc: {formula}.",
            "WFX Smart",
        )
    return read_only_columns


def _write_costing_formulas(ws: Any, last_row: int) -> None:
    columns = {
        label: FORM_COLUMNS.index(label) + 1
        for label in (
            "Cons. Qty.",
            "Waste %",
            "Cons. Qty. Incl. Waste",
            "Rate",
            "Value in (USD)",
        )
    }
    for row in range(2, last_row + 1):
        cons_ref = f"{get_column_letter(columns['Cons. Qty.'])}{row}"
        waste_ref = f"{get_column_letter(columns['Waste %'])}{row}"
        cons_incl_ref = (
            f"{get_column_letter(columns['Cons. Qty. Incl. Waste'])}{row}"
        )
        rate_ref = f"{get_column_letter(columns['Rate'])}{row}"
        cons_incl_cell = ws.cell(
            row=row,
            column=columns["Cons. Qty. Incl. Waste"],
        )
        value_cell = ws.cell(row=row, column=columns["Value in (USD)"])
        cons_incl_cell.value = (
            f'=IF({cons_ref}="","",{cons_ref}*'
            f'(1+IF({waste_ref}="",0,{waste_ref})/100))'
        )
        value_cell.value = (
            f'=IF(OR({rate_ref}="",{cons_incl_ref}=""),"",'
            f"{rate_ref}*{cons_incl_ref})"
        )
        cons_incl_cell.number_format = "0.0000"
        value_cell.number_format = "0.0000"


def _add_dependency_mapping_comments(
    ws: Any,
    document: Mapping[str, Any],
    row_by_item: Mapping[tuple[str, str], int],
) -> None:
    mapping_columns = {
        "colcolordependencymapping": FORM_COLUMNS.index("Color Mapping") + 1,
        "colsizedependencymapping": FORM_COLUMNS.index("Size Mapping") + 1,
    }
    for field in document.get("fields") or ():
        column = mapping_columns.get(_base_field_key(field.get("field_key")))
        if column is None:
            continue
        row = row_by_item.get(
            (
                _text(field.get("section_key")).casefold(),
                _text(field.get("item_key")).casefold(),
            )
        )
        if row is None:
            continue
        choices = _unique_values(field.get("options") or ())
        detail = "\n".join(choices[:100]) or "Chưa scan được option Style."
        cell = ws.cell(row=row, column=column)
        cell.comment = Comment(
            "Mỗi dòng: Material => Style 1 | Style 2\n\n"
            "Các lựa chọn Style đã scan:\n" + detail,
            "WFX Smart",
        )
        cell.alignment = Alignment(vertical="top", wrap_text=True)


def _style_form_sections(
    ws: Any,
    section_rows: Sequence[Sequence[int]],
    visible_column_count: int,
) -> None:
    for rows in section_rows:
        if not rows:
            continue
        for row in rows:
            section_cell = ws.cell(row=row, column=1)
            section_cell.fill = _SUBHEADER_FILL
            section_cell.font = Font(color="0F5660", bold=True)
        for column in range(1, visible_column_count + 1):
            ws.cell(row=rows[0], column=column).border = _SECTION_BORDER


def _style_form_templates(
    ws: Any,
    template_rows: Sequence[tuple[int, str]],
    visible_column_count: int,
    read_only_columns: set[int],
) -> None:
    article_code_column = FORM_COLUMNS.index("Article Code") + 1
    article_name_column = FORM_COLUMNS.index("Article Name") + 1
    for row, section_name in template_rows:
        ws.cell(row=row, column=1).font = Font(
            color="0F5660",
            bold=True,
            italic=True,
        )
        for column in range(2, visible_column_count + 1):
            if column not in read_only_columns:
                ws.cell(row=row, column=column).fill = _TEMPLATE_INPUT_FILL
        section = {"section_key": section_name, "name": section_name}
        is_cost_line = _section_item_type(section) == "cost_line"
        hint = (
            f"Chọn Article Name cho {section_name}; để trống nếu không dùng."
            if is_cost_line
            else (
                f"Dòng thêm Article mới cho {section_name}. "
                "Nhập Article Code hoặc Article Name; để trống nếu không dùng."
            )
        )
        ws.cell(row=row, column=article_code_column).comment = Comment(
            hint,
            "WFX Smart",
        )
        ws.cell(row=row, column=article_name_column).comment = Comment(
            hint,
            "WFX Smart",
        )


def _finish_costing_form(
    ws: Any,
    layout: _CostingFormLayout,
    visible_column_count: int,
) -> None:
    ws.freeze_panes = "E2"
    ws.auto_filter.ref = (
        f"A1:{get_column_letter(visible_column_count)}{max(1, ws.max_row)}"
    )
    ws.row_dimensions[1].height = 34
    for row in range(2, ws.max_row + 1):
        ws.row_dimensions[row].height = 23
    for mapping_label in ("Color Mapping", "Size Mapping"):
        column = FORM_COLUMNS.index(mapping_label) + 1
        for row in layout.row_by_item.values():
            line_count = str(ws.cell(row=row, column=column).value or "").count(
                "\n"
            ) + 1
            if line_count > 1:
                ws.row_dimensions[row].height = max(
                    ws.row_dimensions[row].height or 23,
                    min(120, 16 * line_count),
                )
    for column in range(visible_column_count + 1, len(FORM_COLUMNS) + 1):
        ws.column_dimensions[get_column_letter(column)].hidden = True
    ws.sheet_view.showGridLines = False
    ws.sheet_properties.tabColor = "0F766E"


def _write_costing_form(
    workbook: Workbook,
    document: Mapping[str, Any],
) -> None:
    """Ghi form một sheet với bộ cột cố định và dòng thêm Article sẵn có."""
    ws = workbook.create_sheet(FORM_SHEET)
    _set_header(ws, 1, FORM_COLUMNS)
    visible_column_count = len(FORM_COLUMNS) - len(FORM_TECH_COLUMNS)
    layout = _write_form_data_rows(ws, document)

    form_data_last_row = ws.max_row
    last_row = max(form_data_last_row, 200)
    action_validation = DataValidation(
        type="list",
        formula1='"DELETE"',
        allow_blank=True,
    )
    ws.add_data_validation(action_validation)
    action_validation.add(f"B2:B{last_row}")
    next_lookup_column = _add_form_dropdowns(
        ws,
        document,
        last_row=last_row,
    )
    next_lookup_column = _add_special_article_dropdowns(
        ws,
        document,
        layout,
        lookup_column=next_lookup_column,
    )
    next_lookup_column = _add_material_article_dropdowns(
        ws,
        document,
        layout,
        lookup_column=next_lookup_column,
    )
    _add_item_option_dropdowns(
        ws,
        document,
        layout.row_by_item,
        lookup_column=next_lookup_column,
    )
    for offset, definition in enumerate(
        STANDARD_ITEM_FIELDS,
        len(FORM_BASE_COLUMNS) + 1,
    ):
        ws.cell(row=1, column=offset).comment = Comment(
            f"WFX Field Key: {definition['field_key']}",
            "WFX Smart",
        )
    _finish_sheet(
        ws,
        _form_column_widths(),
        editable_columns=range(2, visible_column_count + 1),
    )
    read_only_columns = _style_formula_columns(ws, form_data_last_row)
    _write_costing_formulas(ws, form_data_last_row)
    _add_dependency_mapping_comments(ws, document, layout.row_by_item)
    _style_form_sections(ws, layout.section_rows, visible_column_count)
    _style_form_templates(
        ws,
        layout.template_rows,
        visible_column_count,
        read_only_columns,
    )
    _finish_costing_form(ws, layout, visible_column_count)


def write_costing_xlsx(document: Mapping[str, Any], path: str | Path) -> Path:
    normalized = workbook_document(document)
    target = _preflight_path(path, must_exist=False)
    if target.suffix.casefold() != ".xlsx":
        raise CostingWorkbookError(
            "COSTING_FILE_TYPE_UNSUPPORTED",
            "Đường dẫn export XLSX phải kết thúc bằng .xlsx.",
        )
    target.parent.mkdir(parents=True, exist_ok=True)
    workbook = Workbook()
    workbook.remove(workbook.active)
    workbook.calculation.calcMode = "auto"
    workbook.calculation.fullCalcOnLoad = True
    workbook.calculation.forceFullCalc = True
    _write_guide(workbook, normalized)
    _write_costing_form(workbook, normalized)
    workbook.save(target)
    return target


def write_costing_file(document: Mapping[str, Any], path: str | Path) -> Path:
    target = _preflight_path(path, must_exist=False)
    return write_costing_xlsx(document, target)
