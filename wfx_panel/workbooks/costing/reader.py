"""Đọc file Costing người dùng nộp lại thành document chuẩn.

Khi user chọn Article Name trong dropdown, lúc đọc phải đồng bộ ngược Article
Code nếu tên chỉ khớp đúng một mã; tên trùng nhiều mã phải báo chọn Code chứ
không được tự đoán."""

from __future__ import annotations

import re
from collections.abc import Iterable, Mapping, Sequence
from dataclasses import dataclass
from dataclasses import field as dataclass_field
from pathlib import Path
from typing import Any

from openpyxl import load_workbook
from openpyxl.utils import column_index_from_string, get_column_letter

from wfx_panel.workbooks.costing.schema import (
    FORM_COLUMNS,
    FORM_SHEET,
    GUIDE_SHEET,
    ITEM_ACTIONS,
    ITEM_TYPES,
    OPTIONAL_FORM_COLUMNS,
    STANDARD_ITEM_FIELDS,
    CostingWorkbookError,
    _clean_cell,
    _form_field_key,
    _normalized_field,
    _normalized_item,
    _order,
    _preflight_path,
    _standard_section_token,
    _text,
    workbook_document,
)


def _worksheet_rows(
    ws: Any,
    *,
    ignored_columns: set[str] | None = None,
) -> Iterable[dict[str, Any]]:
    ignored = ignored_columns or set()
    rows = ws.iter_rows(values_only=True)
    try:
        header = next(rows)
    except StopIteration:
        return
    columns = [_text(value) for value in header]
    for row_index, values in enumerate(rows, 2):
        if not any(value not in (None, "") for value in values):
            continue
        output: dict[str, Any] = {}
        for column_index, key in enumerate(columns):
            if not key or key in ignored:
                continue
            value = values[column_index] if column_index < len(values) else None
            output[key] = _clean_cell(
                value,
                f"{ws.title}!{get_column_letter(column_index + 1)}{row_index}",
            )
        output["__Excel Row"] = row_index
        yield output


def _read_guide_meta(workbook: Any) -> dict[str, Any]:
    if workbook.sheetnames != [GUIDE_SHEET, FORM_SHEET]:
        raise CostingWorkbookError(
            "COSTING_FORMAT_UNSUPPORTED",
            "Workbook Costing phải chỉ có hai sheet: Hướng dẫn và Costing.",
        )
    values = {
        _text(row[0]): _clean_cell(row[1], f"{GUIDE_SHEET}!B{index}")
        for index, row in enumerate(
            workbook[GUIDE_SHEET].iter_rows(values_only=True),
            1,
        )
        if len(row) >= 2 and _text(row[0])
    }
    version = _text(values.get("Phiên bản file")).removeprefix("v")
    return {
        "format_version": version,
        "style_code": _text(values.get("Style Code")),
        "style_name": _text(values.get("Style Name")),
        "cost_sheet_status": _text(values.get("Trạng thái khi tải")),
        "title": "",
        "cost_sheet_type": "Internal Cost Sheets",
        "order_execution_type": "Trading",
        "season": "",
        "template": "FOB",
        "signature": "",
    }


@dataclass
class _CostingFormReadState:
    sections: list[dict[str, Any]] = dataclass_field(default_factory=list)
    items: list[dict[str, Any]] = dataclass_field(default_factory=list)
    fields: list[dict[str, Any]] = dataclass_field(default_factory=list)
    errors: list[str] = dataclass_field(default_factory=list)
    section_keys: set[str] = dataclass_field(default_factory=set)
    item_locations: dict[tuple[str, str], int] = dataclass_field(
        default_factory=dict
    )


@dataclass(frozen=True)
class _FormItemMetadata:
    section_key: str
    item_key: str
    action: str
    item_type: str
    row_index: int


def _missing_form_columns(ws: Any) -> list[str]:
    header = [
        _text(cell.value)
        for cell in next(ws.iter_rows(min_row=1, max_row=1))
    ]
    return [
        column
        for column in FORM_COLUMNS
        if column not in header and column not in OPTIONAL_FORM_COLUMNS
    ]


_ARTICLE_NAME_FORMULA = re.compile(
    r'^=IFERROR\(INDEX\(\$([A-Z]+)\$2:\$\1\$(\d+),'
    r'MATCH\(([A-Z]+)(\d+),\$([A-Z]+)\$2:\$\5\$\2,0\)\),""\)$'
)


_ARTICLE_VALIDATION_RANGE = re.compile(
    r"^=\$([A-Z]+)\$2:\$\1\$(\d+)$"
)


def _resolve_generated_article_name_formulas(ws: Any) -> None:
    """Đổi riêng công thức lookup do app tạo thành text trước khi validate."""
    code_column = FORM_COLUMNS.index("Article Code") + 1
    name_column = FORM_COLUMNS.index("Article Name") + 1
    code_letter = get_column_letter(code_column)
    lookup_cache: dict[tuple[str, str, int], dict[str, str]] = {}
    for row in range(2, ws.max_row + 1):
        cell = ws.cell(row, name_column)
        formula = cell.value
        if not isinstance(formula, str) or not formula.startswith("="):
            continue
        match = _ARTICLE_NAME_FORMULA.fullmatch(formula)
        if match is None:
            continue
        (
            lookup_name_letter,
            last_row_text,
            target_code_letter,
            target_row_text,
            lookup_code_letter,
        ) = match.groups()
        last_row = int(last_row_text)
        if (
            target_code_letter != code_letter
            or int(target_row_text) != row
            or last_row < 2
            or last_row > ws.max_row
            or not ws.column_dimensions[lookup_code_letter].hidden
            or not ws.column_dimensions[lookup_name_letter].hidden
        ):
            continue
        cache_key = (
            lookup_code_letter,
            lookup_name_letter,
            last_row,
        )
        lookup = lookup_cache.get(cache_key)
        if lookup is None:
            lookup = {
                _text(ws[f"{lookup_code_letter}{lookup_row}"].value).casefold():
                _text(ws[f"{lookup_name_letter}{lookup_row}"].value)
                for lookup_row in range(2, last_row + 1)
                if _text(ws[f"{lookup_code_letter}{lookup_row}"].value)
            }
            lookup_cache[cache_key] = lookup
        article_code = _text(ws.cell(row, code_column).value).casefold()
        cell.value = lookup.get(article_code, "")


def _resolve_article_codes_selected_by_name(ws: Any) -> None:
    """Đồng bộ ngược Code khi người dùng chọn Name từ dropdown của app."""
    code_column = FORM_COLUMNS.index("Article Code") + 1
    name_column = FORM_COLUMNS.index("Article Name") + 1
    errors: list[str] = []
    for validation in ws.data_validations.dataValidation:
        formula = str(validation.formula1 or "")
        match = _ARTICLE_VALIDATION_RANGE.fullmatch(formula)
        if match is None:
            continue
        lookup_name_letter, last_row_text = match.groups()
        lookup_name_column = column_index_from_string(lookup_name_letter)
        lookup_code_column = lookup_name_column - 1
        last_row = int(last_row_text)
        if (
            lookup_code_column < 1
            or last_row < 2
            or last_row > ws.max_row
            or not ws.column_dimensions[lookup_name_letter].hidden
            or not ws.column_dimensions[
                get_column_letter(lookup_code_column)
            ].hidden
        ):
            continue
        codes_by_name: dict[str, list[str]] = {}
        names_by_code: dict[str, str] = {}
        for lookup_row in range(2, last_row + 1):
            name = _text(ws.cell(lookup_row, lookup_name_column).value)
            code = _text(ws.cell(lookup_row, lookup_code_column).value)
            if not name or not code:
                continue
            names_by_code.setdefault(code.casefold(), name)
            matching_codes = codes_by_name.setdefault(name.casefold(), [])
            if code.casefold() not in {
                value.casefold() for value in matching_codes
            }:
                matching_codes.append(code)
        for target_range in validation.ranges.ranges:
            if not (
                target_range.min_col <= name_column <= target_range.max_col
            ):
                continue
            for row in range(
                max(2, target_range.min_row),
                min(ws.max_row, target_range.max_row) + 1,
            ):
                name_cell = ws.cell(row, name_column)
                article_name = _text(name_cell.value)
                if not article_name:
                    continue
                if article_name.startswith("="):
                    # Excel/WPS có thể tự viết lại công thức app sinh (thêm @,
                    # _xlfn hoặc đổi cách đặt ngoặc). Chỉ cho phép ở đúng ô
                    # Article Name có validation trỏ vào hai cột lookup ẩn;
                    # không chạy công thức mà thay bằng text từ Code cùng dòng.
                    current_code = _text(ws.cell(row, code_column).value)
                    name_cell.value = names_by_code.get(
                        current_code.casefold(),
                        "",
                    )
                    continue
                matching_codes = codes_by_name.get(article_name.casefold(), [])
                if len(matching_codes) == 1:
                    ws.cell(row, code_column).value = matching_codes[0]
                    continue
                current_code = _text(ws.cell(row, code_column).value)
                if len(matching_codes) > 1 and current_code.casefold() not in {
                    value.casefold() for value in matching_codes
                }:
                    errors.append(
                        f"{FORM_SHEET}!{name_cell.coordinate}: Article Name "
                        f"“{article_name}” trùng {len(matching_codes)} mã; "
                        "hãy chọn Article Code."
                    )
    if errors:
        raise CostingWorkbookError(
            "COSTING_VALIDATION_FAILED",
            "File Costing có Article Name chưa xác định được mã.",
            details=errors[:100],
        )


def _costing_form_rows(ws: Any) -> list[dict[str, Any]]:
    _resolve_generated_article_name_formulas(ws)
    _resolve_article_codes_selected_by_name(ws)
    return list(
        _worksheet_rows(
            ws,
            ignored_columns={
                str(definition["label"])
                for definition in STANDARD_ITEM_FIELDS
                if definition.get("read_only")
            },
        )
    )


def _form_section_key_index(
    rows: Sequence[Mapping[str, Any]],
) -> dict[str, str]:
    return {
        _text(row.get("Section")).casefold(): _text(row.get("__Section Key"))
        for row in rows
        if _text(row.get("Section")) and _text(row.get("__Section Key"))
    }


def _register_form_section(
    state: _CostingFormReadState,
    section_key: str,
    section_name: str,
) -> None:
    normalized_key = section_key.casefold()
    if not section_key or normalized_key in state.section_keys:
        return
    state.section_keys.add(normalized_key)
    state.sections.append(
        {
            "section_key": section_key,
            "name": section_name or section_key,
            "row_order": len(state.sections) + 1,
        }
    )


def _form_row_has_data(row: Mapping[str, Any]) -> bool:
    if _text(row.get("__Item Type")).casefold() == "cost_line":
        # Các dòng mẫu đặc biệt có sẵn Minutes/Curr.; chỉ tạo dòng khi người
        # dùng thật sự chọn tên nhà máy/quy trình/chi phí.
        return bool(_text(row.get("Article Name")))
    if _text(row.get("Article Code")) or _text(row.get("Article Name")):
        return True
    return any(
        row.get(str(definition["label"]), "") not in (None, "")
        for definition in STANDARD_ITEM_FIELDS
    )


def _form_item_key(
    row: Mapping[str, Any],
    section_key: str,
    article_code: str,
    article_name: str,
    row_index: int,
) -> str:
    existing_key = _text(row.get("__Item Key"))
    if existing_key:
        return existing_key
    identity = article_code or article_name or f"row-{row_index}"
    # Adjacent duplicate Articles are valid Splitter requests. Excel row keeps
    # each new row uniquely addressable for planner occurrence matching.
    return f"new:{section_key}:{identity}:row-{row_index}"


def _validate_form_item_row(
    state: _CostingFormReadState,
    metadata: _FormItemMetadata,
) -> None:
    if not metadata.section_key:
        state.errors.append(
            f"{FORM_SHEET}!A{metadata.row_index}: Section không thuộc form chuẩn."
        )
    if metadata.action not in ITEM_ACTIONS:
        state.errors.append(
            f"{FORM_SHEET}!B{metadata.row_index}: "
            f"Action “{metadata.action}” không hợp lệ."
        )
    if metadata.item_type not in ITEM_TYPES:
        item_type_column = FORM_COLUMNS.index("__Item Type") + 1
        state.errors.append(
            f"{FORM_SHEET}!{get_column_letter(item_type_column)}"
            f"{metadata.row_index}: Item Type “{metadata.item_type}” "
            "không hợp lệ."
        )
    composite = (
        metadata.section_key.casefold(),
        metadata.item_key.casefold(),
    )
    if all(composite) and composite in state.item_locations:
        item_key_column = FORM_COLUMNS.index("__Item Key") + 1
        state.errors.append(
            f"{FORM_SHEET}!{get_column_letter(item_key_column)}"
            f"{metadata.row_index}: Item Key trùng dòng "
            f"{state.item_locations[composite]}."
        )
        return
    state.item_locations[composite] = metadata.row_index


def _form_fields_for_item(
    row: Mapping[str, Any],
    section_key: str,
    section_name: str,
    item_key: str,
) -> list[dict[str, Any]]:
    fields: list[dict[str, Any]] = []
    section = {"section_key": section_key, "name": section_name}
    for field_order, definition in enumerate(STANDARD_ITEM_FIELDS):
        if definition.get("read_only"):
            continue
        value = row.get(str(definition["label"]), "")
        if value in (None, ""):
            continue
        field_data = {
            "scope": "item",
            "section_key": section_key,
            "item_key": item_key,
            "field_key": _form_field_key(definition, section),
            "label": definition["label"],
            "value": value,
            "data_type": definition["data_type"],
            "editable": True,
            "required": False,
            "options": [],
            "row_order": field_order,
        }
        fields.append(_normalized_field(field_data, field_order))
        if (
            _standard_section_token(section) == "productioncosts"
            and definition["label"] == "Minutes"
        ):
            fields.append(
                _normalized_field(
                    {
                        **field_data,
                        "field_key": "ProductionHeaderMinutes",
                        "row_order": field_order - 1,
                    },
                    field_order,
                )
            )
    return fields


def _read_costing_form_row(
    row: Mapping[str, Any],
    fallback_row_index: int,
    section_key_by_name: Mapping[str, str],
    state: _CostingFormReadState,
) -> None:
    row_index = _order(row.get("__Excel Row"), fallback_row_index)
    section_name = _text(row.get("Section"))
    section_key = _text(row.get("__Section Key")) or section_key_by_name.get(
        section_name.casefold(),
        "",
    )
    _register_form_section(state, section_key, section_name)
    if not _form_row_has_data(row):
        return
    article_code = _text(row.get("Article Code"))
    article_name = _text(row.get("Article Name"))
    item_key = _form_item_key(
        row,
        section_key,
        article_code,
        article_name,
        row_index,
    )
    action = _text(row.get("Action") or "UPSERT").upper()
    item_type = _text(row.get("__Item Type") or "article").casefold()
    item_fields = _form_fields_for_item(
        row,
        section_key,
        section_name,
        item_key,
    )
    if (
        action == "UPSERT"
        and _text(row.get("__Item Key"))
        and not item_fields
    ):
        # Bản export cũ có thể gắn subtotal chỉ đọc vào một dòng ``>>`` ẩn.
        # Sau khi bỏ field công thức, dòng này chỉ còn identity và không có gì
        # để apply; bỏ qua để file cũ không tạo split/Purchase Officer giả.
        return
    _validate_form_item_row(
        state,
        _FormItemMetadata(
            section_key=section_key,
            item_key=item_key,
            action=action,
            item_type=item_type,
            row_index=row_index,
        ),
    )
    state.items.append(
        _normalized_item(
            {
                "section_key": section_key,
                "section_name": section_name,
                "item_key": item_key,
                "row_order": _order(row.get("__Row Order"), row_index),
                "action": action,
                "item_type": item_type,
                "article_code": article_code,
                "article_name": article_name,
            },
            row_index,
        )
    )
    state.fields.extend(item_fields)


def _read_costing_form(
    workbook: Any,
) -> tuple[list[dict[str, Any]], list[dict[str, Any]], list[dict[str, Any]]]:
    ws = workbook[FORM_SHEET]
    missing = _missing_form_columns(ws)
    if missing:
        raise CostingWorkbookError(
            "COSTING_FORMAT_UNSUPPORTED",
            "Sheet Costing thiếu cột chuẩn của WFX Smart.",
            details=[
                f"{FORM_SHEET}!hàng 1: thiếu cột “{column}”."
                for column in missing
            ],
        )
    rows = _costing_form_rows(ws)
    section_key_by_name = _form_section_key_index(rows)
    state = _CostingFormReadState()
    for fallback_row_index, row in enumerate(rows, 2):
        _read_costing_form_row(
            row,
            fallback_row_index,
            section_key_by_name,
            state,
        )
    if state.errors:
        raise CostingWorkbookError(
            "COSTING_VALIDATION_FAILED",
            "File Costing có dữ liệu chưa hợp lệ.",
            details=state.errors[:100],
        )
    return state.sections, state.items, state.fields


def read_costing_xlsx(path: str | Path) -> dict[str, Any]:
    target = _preflight_path(path, must_exist=True)
    try:
        workbook = load_workbook(
            target,
            read_only=False,
            data_only=False,
            keep_links=False,
        )
    except CostingWorkbookError:
        raise
    except Exception as error:
        raise CostingWorkbookError(
            "COSTING_VALIDATION_FAILED",
            "Không đọc được workbook Costing.",
        ) from error
    meta = _read_guide_meta(workbook)
    if not _text(meta.get("style_code")):
        raise CostingWorkbookError(
            "COSTING_VALIDATION_FAILED",
            "File Costing có dữ liệu chưa hợp lệ.",
            details=[f"{GUIDE_SHEET}!B2: thiếu Style Code."],
        )
    sections, items, fields = _read_costing_form(workbook)
    document = {
        **meta,
        "fields": fields,
        "sections": sections,
        "items": items,
    }
    return workbook_document(document)


def read_costing_file(path: str | Path) -> dict[str, Any]:
    target = _preflight_path(path, must_exist=True)
    return read_costing_xlsx(target)


def costing_file_summary(
    document: Mapping[str, Any],
    path: str | Path,
) -> dict[str, Any]:
    normalized = workbook_document(document)
    target = Path(path)
    return {
        "file_name": target.name,
        "file_format": target.suffix.casefold().lstrip("."),
        "style_code": normalized["style_code"],
        "section_count": len(normalized["sections"]),
        "item_count": len(normalized["items"]),
        "field_count": len(normalized["fields"]),
    }
