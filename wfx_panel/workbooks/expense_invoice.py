"""Đọc mẫu Expense Invoice; chỉ dữ liệu đã validate mới được gửi sang WFX."""

from __future__ import annotations

import json
import unicodedata
from datetime import date, datetime
from decimal import Decimal, InvalidOperation
from pathlib import Path
from zipfile import BadZipFile, ZipFile

from openpyxl import load_workbook
from openpyxl.workbook.defined_name import DefinedName
from openpyxl.worksheet.datavalidation import DataValidation

from wfx_panel.paths import RESOURCE_DIR

SHEET = "Expense Invoice"
COLUMNS = (
    "Supplier",
    "Invoice No.",
    "Invoice Date",
    "GL Date",
    "Invoice Currency",
    "Payment Term",
    "Due Date",
    "Ledger",
    "Qty",
    "Rate",
    "Comments",
    "Brand",
    "Team",
)
KEYS = (
    "supplier",
    "invoice_no",
    "invoice_date",
    "gl_date",
    "currency",
    "payment_term",
    "due_date",
    "ledger",
    "qty",
    "rate",
    "comments",
    "brand",
    "team",
)
HEADER_KEYS = (*KEYS[:7], *KEYS[10:])
MAX_ROWS = 500
TEMPLATE = RESOURCE_DIR / "wfx_panel/assets/templates/Expense-Invoice.xlsx"
# Dropdown ban đầu quét từ WFX lúc phát hành (scripts/seed_expense_options.py);
# app thay bằng danh sách thật của từng tài khoản/Division sau lượt điền đầu.
OPTIONS_SEED = RESOURCE_DIR / "wfx_panel/assets/templates/Expense-Invoice-options.json"
OPTION_KEYS = ("supplier", "currency", "payment_term", "ledger", "brand", "team")


class ExpenseWorkbookError(ValueError):
    pass


def load_seed_options() -> dict:
    try:
        options = json.loads(OPTIONS_SEED.read_text(encoding="utf-8"))["options"]
    except (OSError, ValueError, KeyError, TypeError):
        return {}
    return options if isinstance(options, dict) else {}


def text(value: object) -> str:
    if value is None:
        return ""
    if isinstance(value, float) and value.is_integer():
        return str(int(value))
    return unicodedata.normalize("NFC", str(value)).strip()


def _day(value: object) -> str:
    if isinstance(value, datetime):
        value = value.date()
    if isinstance(value, date):
        return value.isoformat()
    for fmt in ("%Y-%m-%d", "%d/%m/%Y", "%d %b %Y"):
        try:
            return datetime.strptime(text(value), fmt).date().isoformat()
        except ValueError:
            continue
    raise ValueError("ngày phải là ô ngày Excel, yyyy-mm-dd hoặc dd/mm/yyyy")


def number(value: object, *, qty: bool = False) -> Decimal:
    if isinstance(value, bool):
        raise ValueError("phải là số")
    try:
        result = Decimal(text(value))
    except InvalidOperation as error:
        raise ValueError("phải là số, không có dấu phân cách hàng nghìn") from error
    if not result.is_finite() or result < 0 or (qty and result == 0):
        raise ValueError("Qty phải lớn hơn 0; Rate không được âm")
    places = 2 if qty else 4
    if result > Decimal("9999999999") or result != result.quantize(
        Decimal(10) ** -places
    ):
        raise ValueError(f"tối đa 10 chữ số phần nguyên và {places} chữ số thập phân")
    return result


def save_template(target: Path, options: dict | None = None) -> None:
    """Bổ sung danh sách WFX vào mẫu đã đóng gói bằng thư viện của ứng dụng.

    Thiếu danh sách nào thì cột đó để nhập tự do; mẫu vẫn dùng được."""
    workbook = load_workbook(TEMPLATE)
    try:
        sheet = workbook[SHEET]
        lists = workbook.create_sheet("_Lists")
        fields = {
            "supplier": "A",
            "currency": "E",
            "payment_term": "F",
            "ledger": "H",
            "brand": "L",
            "team": "M",
        }
        for index, (key, column) in enumerate(fields.items(), 1):
            values = list(
                dict.fromkeys(text(v) for v in (options or {}).get(key, []) if text(v))
            )
            if not values:
                continue
            letter = chr(64 + index)
            lists.cell(1, index, key)
            for row, value in enumerate(values, 2):
                lists.cell(row, index, value).data_type = "s"
            name = f"Expense_{key}"
            workbook.defined_names.add(
                DefinedName(
                    name, attr_text=f"'_Lists'!${letter}$2:${letter}${len(values) + 1}"
                )
            )
            validation = DataValidation(
                type="list", formula1=f"={name}", allow_blank=True
            )
            validation.errorTitle = "Giá trị chưa có trong danh sách"
            validation.error = (
                "Danh sách có thể chưa có Supplier/Ledger mới. Chọn Yes nếu đúng "
                "tên trên WFX; app vẫn đối chiếu lại khi điền."
            )
            # Danh sách đi kèm app có thể cũ hơn WFX: chỉ cảnh báo, không chặn.
            validation.errorStyle = "warning"
            validation.showErrorMessage = True
            sheet.add_data_validation(validation)
            validation.add(f"{column}2:{column}{MAX_ROWS + 1}")
        lists.sheet_state = "veryHidden"
        workbook.save(target)
    finally:
        workbook.close()


def read_expense_workbook(source: Path) -> dict:
    try:
        if source.suffix.casefold() != ".xlsx":
            raise ExpenseWorkbookError("Chỉ hỗ trợ mẫu .xlsx của WFX Smart.")
        if source.stat().st_size > 5 * 1024 * 1024:
            raise ExpenseWorkbookError("File vượt quá giới hạn 5 MB.")
        with ZipFile(source) as archive:
            entries = archive.infolist()
            if (
                len(entries) > 5000
                or sum(e.file_size for e in entries) > 32 * 1024 * 1024
            ):
                raise ExpenseWorkbookError("Nội dung file Excel quá lớn.")
            if any(e.filename.casefold().endswith("vbaproject.bin") for e in entries):
                raise ExpenseWorkbookError("File không được chứa macro.")
        workbook = load_workbook(
            source, read_only=True, data_only=False, keep_links=False
        )
    except ExpenseWorkbookError:
        raise
    except (OSError, BadZipFile, ValueError, KeyError) as error:
        raise ExpenseWorkbookError("Không đọc được file Excel đã chọn.") from error
    try:
        if SHEET not in workbook.sheetnames:
            raise ExpenseWorkbookError(f"Không thấy sheet {SHEET}; hãy tải lại mẫu.")
        sheet = workbook[SHEET]
        if sheet.max_row > MAX_ROWS + 1 or sheet.max_column > len(COLUMNS):
            raise ExpenseWorkbookError("Mẫu có tối đa 500 dòng dữ liệu và đúng 13 cột.")
        records = sheet.iter_rows(max_col=len(COLUMNS))
        header = next(records)
        if tuple(text(c.value) for c in header) != COLUMNS:
            raise ExpenseWorkbookError(
                "Tên hoặc thứ tự 13 cột không đúng mẫu Expense Invoice."
            )
        shared: dict[str, str] = {}
        lines: list[dict] = []
        errors: list[str] = []
        for row_no, cells in enumerate(records, 2):
            if not any(text(c.value) for c in cells):
                continue
            if any(c.data_type in {"f", "e"} for c in cells):
                errors.append(
                    f"Dòng {row_no}: chỉ nhập giá trị, không dùng công thức/ô lỗi."
                )
                continue
            raw = dict(zip(KEYS, (c.value for c in cells), strict=True))
            row: dict[str, str] = {}
            for key, value in raw.items():
                column = COLUMNS[KEYS.index(key)]
                try:
                    if key in HEADER_KEYS and not text(value) and shared:
                        row[key] = shared[key]
                    elif key in {"invoice_date", "gl_date", "due_date"}:
                        row[key] = _day(value)
                    elif key in {"qty", "rate"}:
                        row[key] = format(number(value, qty=key == "qty"), "f")
                    else:
                        row[key] = text(value)
                    if key in {*KEYS[:8]} and not row[key]:
                        raise ValueError("không được để trống ở dòng đầu")
                    if len(row[key]) > (2000 if key == "comments" else 255):
                        raise ValueError("nội dung quá dài")
                except (ValueError, InvalidOperation) as error:
                    errors.append(f"Dòng {row_no} · {column}: {error}.")
            if len(row) != len(KEYS):
                continue
            if not shared:
                shared = {key: row[key] for key in HEADER_KEYS}
            elif any(row[key] != shared[key] for key in HEADER_KEYS):
                errors.append(
                    f"Dòng {row_no}: mỗi file chỉ chứa một hóa đơn; thông tin chung phải giống dòng đầu."
                )
            lines.append({key: row[key] for key in ("ledger", "qty", "rate")})
        if errors:
            raise ExpenseWorkbookError("\n".join(errors[:15]))
        if not lines:
            raise ExpenseWorkbookError("File chưa có dòng chi phí nào.")
        return {
            "header": shared,
            "lines": lines,
            "total": format(
                sum(Decimal(r["qty"]) * Decimal(r["rate"]) for r in lines), "f"
            ),
        }
    finally:
        workbook.close()
