"""Điểm vào: ghép Buyer Invoice và Packing List thành một workbook."""

from __future__ import annotations

import xml.etree.ElementTree as ET
from pathlib import Path
from zipfile import BadZipFile

from openpyxl import load_workbook

from wfx_panel.workbooks.asn.layout import (
    _fit_packing_measurement_columns,
    _fit_reports_to_a4,
    _fit_wrapped_report_rows,
)
from wfx_panel.workbooks.asn.ooxml import ASNWorkbookError, _merge_packages
from wfx_panel.workbooks.asn.packing import (
    _JL_BUYER,
    _TRUEWERK_BUYER,
    _merge_jl_packing_measurements,
    _merge_truewerk_packing_measurements,
)


def merge_sale_asn_reports(
    packing_list_path: str | Path,
    buyer_invoice_path: str | Path,
    output_path: str | Path,
    *,
    invoice_no: str = "",
    buyer_name: str = "",
) -> Path:
    """Ghép report ở cấp OOXML để giữ nguyên khung của cả Invoice và PKL."""
    packing_path = Path(packing_list_path).expanduser().resolve()
    buyer_path = Path(buyer_invoice_path).expanduser().resolve()
    target = Path(output_path).expanduser().resolve()
    if target.suffix.casefold() != ".xlsx":
        raise ASNWorkbookError("File Sale ASN phải có đuôi .xlsx.")
    for source in (packing_path, buyer_path):
        if not source.is_file() or source.stat().st_size <= 0:
            raise ASNWorkbookError(f"Không đọc được report: {source.name}.")
    invoice_label = str(invoice_no or target.stem).strip() or "Invoice"
    buyer_key = " ".join(str(buyer_name or "").split()).casefold()
    target.parent.mkdir(parents=True, exist_ok=True)
    try:
        _merge_packages(packing_path, buyer_path, target, invoice_label)
        if buyer_key == _JL_BUYER:
            _merge_jl_packing_measurements(target)
        elif buyer_key == _TRUEWERK_BUYER:
            _merge_truewerk_packing_measurements(target)
        _fit_packing_measurement_columns(target)
        _fit_wrapped_report_rows(target)
        _fit_reports_to_a4(target)
        verified = load_workbook(target, read_only=True, data_only=False)
        verified.close()
    except ASNWorkbookError:
        target.unlink(missing_ok=True)
        raise
    except (BadZipFile, KeyError, ET.ParseError) as error:
        target.unlink(missing_ok=True)
        raise ASNWorkbookError(
            f"Report WFX không phải workbook Excel hợp lệ: {error}"
        ) from error
    except Exception as error:
        target.unlink(missing_ok=True)
        raise ASNWorkbookError(f"Không ghép được hai report Sale ASN: {error}") from error
    return target


def sale_asn_sheet_names(path: str | Path) -> list[str]:
    """Đọc tên sheet thực tế sau khi ghép để UI không báo tên giả định."""
    source = Path(path).expanduser().resolve()
    try:
        workbook = load_workbook(source, read_only=True, data_only=False)
    except Exception as error:
        raise ASNWorkbookError(f"Không đọc được tên sheet Sale ASN: {error}") from error
    try:
        return list(workbook.sheetnames)
    finally:
        workbook.close()
